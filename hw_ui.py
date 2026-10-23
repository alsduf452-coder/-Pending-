"""Streamlit 설비 H/W 이력관리 화면."""
from datetime import datetime
from io import BytesIO
import json
import math
from pathlib import Path
import sqlite3
import tempfile
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED
from zoneinfo import ZoneInfo
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
import hw_store as store
from pending_data import safe_export


def card_menu(label,options,key,horizontal=False):
    if st.session_state.get(key) not in options:st.session_state[key]=options[0]
    def select(value):st.session_state[key]=value
    with st.container(key=f'cards_{key}'):
        st.caption(label)
        targets=st.columns(len(options),gap='small') if horizontal else [st]*len(options)
        for target,option in zip(targets,options):
            target.button(option,key=f'nav_{key}_{option}',type='primary' if st.session_state[key]==option else 'secondary',use_container_width=True,on_click=select,args=(option,))
    return st.session_state[key]


def today(): return datetime.now(ZoneInfo('Asia/Seoul')).date()


def text(value): return '' if value is None else str(value)


def table(rows):
    if rows:st.dataframe(pd.DataFrame(rows),hide_index=True,use_container_width=True)
    else:st.info('등록된 기록이 없습니다.')


def show_audit(path,identifier=None):
    rows=store.audit_records(path,identifier)
    if not rows:st.info('변경 이력이 없습니다.');return
    readable=[]
    for r in rows:
        before=json.loads(r['before_json']) if r['before_json'] else {}
        after=json.loads(r['after_json']) if r['after_json'] else {}
        fields=[key for key in set(before)|set(after) if before.get(key)!=after.get(key)]
        stamp=pd.to_datetime(r['timestamp'],utc=True).tz_convert('Asia/Seoul').strftime('%Y-%m-%d %H:%M:%S KST')
        for field in sorted(fields):
            readable.append({'시각':stamp,'변경자':r['actor'],'작업':r['action'],'종류':store.KINDS[r['kind']],'항목 ID':r['entity_id'],'필드':field,'변경 전':json.dumps(before.get(field),ensure_ascii=False),'변경 후':json.dumps(after.get(field),ensure_ascii=False)})
    df=pd.DataFrame(readable)
    st.dataframe(df,hide_index=True,use_container_width=True)
    st.download_button('변경 이력 CSV',safe_export(df).to_csv(index=False).encode('utf-8-sig'),'hw_audit.csv','text/csv',key=f'audit_{identifier or "all"}')


def filter_rows(rows,equipment,chamber,query=''):
    result=[r for r in rows if r.get('equipmentId')==equipment and (chamber=='전체' or r.get('chamberId')==chamber)]
    for word in query.split(): result=[r for r in result if word.casefold() in json.dumps(r,ensure_ascii=False).casefold()]
    return result


def pending_picker(context,equipment,current=None):
    links={}
    if context:
        frame=context['frame']
        eq=equipment.removeprefix('A4')
        for _,row in frame[frame['설비'].str.contains(eq,regex=False)].iterrows():
            links[row['ID']]={'source':context['source'],'id':row['ID'],'title':str(row['내용'])[:160]}
    if current:
        existing_key=current['id']
        # 다른 Excel에 저장된 링크는 별도 키로 유지하여 같은 행 번호로 바뀌지 않게 합니다.
        if not context or current.get('source')!=context['source']:existing_key='기존 연결 유지'
        links[existing_key]=current
    options=['연결 없음']+list(links)
    default=next((k for k,v in links.items() if v==current),'연결 없음')
    selected=st.selectbox('관련 Pending',options,index=options.index(default),format_func=lambda k:k if k=='연결 없음' else f"{links[k]['id']} · {links[k].get('title','')[:60]}")
    return links.get(selected)


def show_pending_link(link,context):
    if not link: return
    st.caption(f"관련 Pending: {link['id']} · {link.get('title','')}")
    if context and link.get('source')==context['source']:
        matched=context['frame'][context['frame']['ID']==link['id']]
        if not matched.empty:
            st.write('현재 Pending 상태:',matched.iloc[0]['상태'])
            def back():
                st.session_state['main_page']='Pending 현황'
                st.session_state['keyword']=link['id']
                st.session_state.pop('count_selection',None)
                st.session_state.pop('list_fields',None)
            st.button('연결된 Pending 조회',on_click=back,key='hw_back_pending')
            st.caption('사이드바 필터로 가려지면 조회 범위·설비·분야·상태를 전체로 선택해주세요.')
            return
    st.warning('현재 선택한 Excel에서 연결 대상을 찾지 못했습니다. 연결과 H/W 기록은 보존됩니다.')


def attachments(path,item):
    for i,attachment in enumerate(item.get('attachments') or []):
        if not isinstance(attachment,dict):continue
        st.caption(attachment.get('name','첨부파일'))
        blob=store.get_blob(path,attachment['blobId']) if attachment.get('blobId') else None
        if blob:
            if blob['mime'] in ['image/png','image/jpeg']:
                try:st.image(blob['data'],caption=blob['name'],width=500)
                except Exception:st.info('이미지 미리보기를 표시하지 못했습니다. 다운로드해 확인해주세요.')
            st.download_button('첨부파일 다운로드',blob['data'],blob['name'],blob['mime'],key=f'blob_{item["id"]}_{i}')
        else:
            st.info('백업에는 파일 경로만 있고 실제 파일이 없습니다. 필요한 원본 파일을 작업 수정에서 첨부해주세요.')


def work_form(path,equipment,chambers,context,current=None):
    current=current or {};identifier=current.get('id') or 'new'
    version=current.get('_version')
    with st.form(f'work_form_{identifier}_{version}'):
        ch=st.selectbox('작업 챔버',chambers,index=chambers.index(current['chamberId']) if current.get('chamberId') in chambers else 0,disabled=bool(current))
        work_date=st.text_input('작업일 * · YYYY-MM-DD',value=current.get('workDate',today().isoformat()))
        work_time=st.text_input('작업 시각 · HH:MM',value=(current.get('workTimestamp') or '')[11:16])
        work_order=st.number_input('같은 날 작업 순서',min_value=1,step=1,value=int(current.get('workOrder') or 1))
        types=['초기 셋업','부품 교체','교체/PM','점검','Tuning','CSD 조정','재측정','기타']
        if current.get('workType') and current['workType'] not in types:types.append(current['workType'])
        work_type=st.selectbox('작업 유형',types,index=types.index(current.get('workType') or '점검'))
        parts=store.PARTS.copy()
        if current.get('targetPart') and current['targetPart'] not in parts:parts.append(current['targetPart'])
        part=st.selectbox('대상 부품',parts,index=parts.index(current.get('targetPart') or '기타 부품'))
        before=st.text_input('교체 전 S/N',value=current.get('beforeSn',''))
        after=st.text_input('교체 후 S/N',value=current.get('afterSn',''))
        reason=st.text_area('작업 사유',value=current.get('reason',''))
        details=st.text_area('작업 내용 *',value=current.get('details',''),max_chars=20000)
        statuses=['정상완료','조치필요','모니터링중','이슈발생']
        status=st.selectbox('작업 결과',statuses,index=statuses.index(current.get('issueStatus','정상완료')) if current.get('issueStatus','정상완료') in statuses else 0)
        worker=st.text_input('작업자',value=current.get('worker',''))
        prefill=st.session_state.get('pending_hw_prefill')
        if prefill and prefill.get('equipmentId') not in [None,equipment]:prefill=None
        pending=pending_picker(context,equipment,current.get('pendingLink') or prefill)
        photos=st.file_uploader('첨부 · PNG/JPG/PDF, 개당 10MB',type=['png','jpg','jpeg','pdf'],accept_multiple_files=True,key=f'work_photos_{identifier}_{version}')
        actor=st.text_input('등록·변경자 *',max_chars=200)
        if st.form_submit_button('작업 이력 저장'):
            try:
                if work_time and not __import__('re').fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d',work_time):raise ValueError('시각은 HH:MM 형식으로 입력해주세요.')
                item={**store.snapshot_record(current),'id':current.get('id') or 'WORK-'+uuid4().hex,'equipmentId':equipment,'chamberId':ch,'workDate':work_date,'workTimestamp':work_date+(' '+work_time if work_time else ''),'workOrder':work_order,'workType':work_type,'targetPart':part,'beforeSn':before.strip(),'afterSn':after.strip(),'reason':reason,'details':details.strip(),'issueStatus':status,'worker':worker,'pendingLink':pending}
                uploads=[{'name':f.name,'data':f.getvalue(),'mime':'application/pdf' if f.name.lower().endswith('.pdf') else 'image/png' if f.name.lower().endswith('.png') else 'image/jpeg'} for f in photos]
                store.save_record(path,'workLogs',item,actor,version,uploads)
                st.session_state.pop('pending_hw_prefill',None)
                st.session_state['hw_saved']='작업 이력을 저장했습니다. Pending 완료 상태는 별도로 처리해주세요.'
                st.rerun()
            except Exception as exc:st.error(f'저장 실패: {exc}')


def measurement_form(path,kind,equipment,chambers,logs,current=None,default_type=None):
    current=current or {};spacing=kind=='spacingMeasurements';identifier=current.get('id','new');version=current.get('_version')
    with st.form(f'measure_{kind}_{identifier}_{version}'):
        ch=st.selectbox('측정 챔버',chambers,index=chambers.index(current['chamberId']) if current.get('chamberId') in chambers else 0,disabled=bool(current))
        date=st.text_input('측정일 * · YYYY-MM-DD',value=current.get('measuredAt',today().isoformat()))
        timing_options=['초기 셋업','작업 전','작업 후','정기 점검']
        timing=st.selectbox('측정 시점',timing_options,index=timing_options.index(current.get('timing','작업 후')) if current.get('timing','작업 후') in timing_options else 0)
        title=st.text_input('작업 제목',value=current.get('workTitle',''))
        part_type=st.selectbox('Flatness 종류',['Diffuser','Susceptor'],index=1 if (current.get('type') or default_type)=='Susceptor' else 0) if not spacing else None
        sn=st.text_input('측정 부품 S/N',value=current.get('partSnAtMeasurement',current.get('partSn','')))
        diffuser=st.text_input('Diffuser S/N',value=current.get('diffuserSn',''))
        susceptor=st.text_input('Susceptor S/N',value=current.get('susceptorSn',''))
        choices={r['id']:r for r in logs if r['equipmentId']==equipment}
        if current.get('relatedWorkLogId') and current['relatedWorkLogId'] not in choices:
            choices[current['relatedWorkLogId']]={'workDate':'기존 연결','details':''}
        options=['연결 없음']+list(choices)
        related=st.selectbox('관련 작업 이력',options,index=options.index(current.get('relatedWorkLogId')) if current.get('relatedWorkLogId') in options else 0,format_func=lambda x:x if x=='연결 없음' else f"{x} · {choices[x]['workDate']}")
        st.caption('단위 mils · 빈칸은 미측정으로 유지합니다. 0은 실제 측정값입니다.')
        values=current.get('values',[None]*9) if spacing else current.get('grid',[[None]*7 for _ in range(5)])
        df=pd.DataFrame([values],columns=[f'P{i}' for i in range(1,10)]) if spacing else pd.DataFrame(values,columns=[f'C{i}' for i in range(1,8)])
        df=df.astype(float)
        editor=st.data_editor(df,num_rows='fixed',use_container_width=True,key=f'grid_{kind}_{identifier}_{version}')
        notes=st.text_area('비고',value=current.get('notes',''))
        actor=st.text_input('등록·변경자 *',max_chars=200)
        if st.form_submit_button('측정 기록 저장'):
            try:
                def number(v):return None if pd.isna(v) else float(v)
                grid=[[number(v) for v in row] for row in editor.to_numpy().tolist()]
                item={**store.snapshot_record(current),'id':current.get('id') or 'MEASURE-'+uuid4().hex,'equipmentId':equipment,'chamberId':ch,'measuredAt':date,'timing':timing,'workTitle':title,'diffuserSn':diffuser,'susceptorSn':susceptor,'relatedWorkLogId':None if related=='연결 없음' else related,'unit':'mils','notes':notes,'gaugeTool':'MCT-2009LT' if spacing else 'Laser-L730'}
                if related!='연결 없음' and choices[related].get('chamberId') not in [None,ch]:raise ValueError('관련 작업 이력의 챔버가 다릅니다.')
                if spacing:item.update(values=grid[0],partSnAtMeasurement=sn)
                else:item.update(grid=grid,type=part_type,partSn=sn,rows=5,cols=7,maxMinSpan=store.flatness_span(grid))
                store.save_record(path,kind,item,actor,version)
                if not spacing:st.session_state[f'flatness_type_{equipment}']=part_type
                st.session_state['hw_saved']='측정 기록을 저장했습니다.';st.rerun()
            except Exception as exc:st.error(f'저장 실패: {exc}')


def measurement_view(path,kind,equipment,chamber,chambers,all_logs):
    spacing=kind=='spacingMeasurements'
    st.markdown("""<style>
.spacing-values {display:grid;grid-template-columns:repeat(3,minmax(0,1fr));border:1px solid #94a3b8;border-radius:9px;overflow:hidden;background:#f1f5f9}
.spacing-cell {min-height:135px;padding:12px;border:1px solid #cbd5e1;display:flex;flex-direction:column;gap:7px;color:#172b4d}
.spacing-cell strong {font-size:27px;line-height:1.25;overflow-wrap:anywhere}.spacing-cell small {font-size:12px}.spacing-cell span {font-size:11px;color:#475569}
.spacing-center {background:#e0f2fe;border:2px solid #0891b2}.measurement-svg {width:100%;height:auto;display:block;background:#f1f5f9;border:1px solid #94a3b8;border-radius:10px}
.flatness-values {width:100%;border-collapse:collapse;table-layout:fixed;background:#f1f5f9;color:#172b4d}.flatness-values th,.flatness-values td {border:1px solid #94a3b8;padding:10px 3px;text-align:center}.flatness-values th {background:#cbd5e1;font-size:12px}.flatness-values td small {display:block;color:#64748b;font-size:9px;text-align:left}.flatness-values td strong {display:block;background:#e2e8f0;border-radius:4px;padding:7px 0;font-size:15px}.flatness-values th small {display:block;font-size:10px;font-weight:normal}
.measure-legend {height:15px;border-radius:9px;border:1px solid #64748b;margin:10px 0}
</style>""",unsafe_allow_html=True)
    scoped=filter_rows(store.records(path,kind),equipment,chamber)
    flatness_type=None if spacing else flatness_type_cards(equipment,scoped)
    query=st.text_input('측정 기록 검색',key=f'search_{kind}')
    rows=filter_rows(scoped,equipment,chamber,query)
    if not spacing:rows=[row for row in rows if row['type']==flatness_type]
    with st.expander('＋ 새 측정 등록'):
        measurement_form(path,kind,equipment,chambers,all_logs,default_type=flatness_type)
    if not rows:
        st.info(f"해당 조건의 {'Spacing' if spacing else flatness_type+' Flatness'} 측정 기록이 없습니다. 다른 카드·챔버·검색 조건을 확인해주세요.")
        return
    indexed={r['id']:r for r in rows}
    suffix='' if spacing else f'_{flatness_type}'
    selected_key=f'measure_selected_{kind}_{equipment}{suffix}'
    selected=st.session_state.get(selected_key)
    if selected not in indexed:
        selected=max(rows,key=lambda row:(row['measuredAt'],row.get('originalOrder') or 0,row['id']))['id']
        st.session_state[selected_key]=selected
    r=indexed[selected]
    st.subheader(f"{r['chamberId']} · {'Spacing' if spacing else r['type']+' Flatness'} · {r['measuredAt']}")
    st.caption('챔버별 측정 이력의 보기·수정 버튼으로 다른 기록을 열 수 있습니다.')
    st.write('측정기:',r.get('gaugeTool',''),' / 단위:',r.get('unit',''))
    st.caption(f"S/N: {r.get('partSnAtMeasurement',r.get('partSn',''))}")
    st.caption('측정 구분: '+str(r.get('timing') or '미기재'))
    if r.get('notes'):st.caption('비고: '+r['notes'])
    if r.get('isAmbiguous'):st.warning(r.get('ambiguityReason') or '원본에서 검토 대상으로 표시된 측정입니다. 값을 자동 보정하지 않았습니다.')
    source=r.get('source',{})
    if source:st.caption(f"원문 출처: {source.get('fileName','')} / {source.get('sheetName','')} / {source.get('cellRange','')}")
    measurement_history(rows,kind,equipment)
    mode_key=f"measurement_mode_{kind}_{equipment}_{r['chamberId']}_{flatness_type or 'Spacing'}"
    view_mode=st.radio('측정 보기 방식',['단일 상세','작업 전후 비교'],horizontal=True,key=mode_key)
    if view_mode=='작업 전후 비교':
        comparison_view(rows,r,kind,equipment)
        measurement_export(rows,kind,spacing)
        return
    titles=st.columns(2,gap='large')
    titles[0].subheader('Spacing 실측 데이터' if spacing else f'{flatness_type} Flatness 실측 데이터')
    titles[1].subheader('Spacing 측정 위치도' if spacing else f'{flatness_type} Flatness 측정 위치도')
    orientation='▲ G/V 측: ①⑥② / 중앙: ⑨⑤⑦ / ▼ Window 측: ④⑧③' if spacing else '5행 × 7열 · Col 1: Gate Valve 측 / Col 7: Window 측'
    for col in titles:col.caption(orientation)
    left,right=st.columns(2,gap='large')
    left_panel=left.empty();right_panel=right.empty()
    editor_key=f"preview_{kind}_{r['id']}_{r['_version']}"
    original_grid=[[r['values'][p-1] for p in positions] for positions in store.POINT_LAYOUT] if spacing else r['grid']
    with st.expander('수치 직접 편집 · 변경값은 위 위치도에 즉시 미리보기',expanded=st.session_state.get('measurement_edit_requested')==r['id']):
        data=pd.DataFrame(original_grid,columns=['좌','중앙','우'] if spacing else [f'Col {i}' for i in range(1,8)],index=['G/V 측','중앙','Window 측'] if spacing else [f'Row {i}' for i in range(1,6)]).astype(float)
        edited=st.data_editor(data,num_rows='fixed',use_container_width=True,key=editor_key)
    draft_grid=[[None if pd.isna(v) else float(v) for v in row] for row in edited.to_numpy().tolist()]
    if spacing:
        draft_values=[None]*9
        for ri,positions in enumerate(store.POINT_LAYOUT):
            for ci,point in enumerate(positions):draft_values[point-1]=draft_grid[ri][ci]
    left_panel.markdown(spacing_table_html(draft_values) if spacing else flatness_table_html(draft_grid),unsafe_allow_html=True)
    right_panel.markdown(spacing_diagram_svg(draft_values) if spacing else flatness_diagram_svg(draft_grid),unsafe_allow_html=True)
    stats=measurement_stats(draft_grid)
    if not spacing:
        legend='#64748b' if stats['최소값']==stats['최댓값'] else 'linear-gradient(90deg,#2563eb,#ef4444)'
        st.markdown(f'<div class="measure-legend" style="background:{legend}"></div>',unsafe_allow_html=True)
        st.caption(f"낮은 값 · 파랑: {measurement_number(stats['최소값'])} mils ↔ 높은 값 · 빨강: {measurement_number(stats['최댓값'])} mils · 회색/—: 미측정")
        if stats['최소값'] is not None and stats['최소값']==stats['최댓값']:st.caption('모든 실측값이 같아 도면은 중립 회색으로 표시합니다.')
    table([{'최소값(mils)':stats['최소값'],'최댓값(mils)':stats['최댓값'],'Max−Min(mils)':stats['Max−Min'],'측정 개수':f"{stats['측정 수']}/{9 if spacing else 35}"}])
    st.caption('왼쪽 실측 데이터와 오른쪽 위치도는 같은 행·열 순서의 값을 표시합니다. 빈칸은 미측정, 0은 실측값입니다. 편집 후 저장해야 DB가 변경됩니다.')
    if spacing:st.caption('885 mils는 AI Studio 원본의 시각 비교 기준입니다. 규격이나 양불 판정 기준이 아닙니다.')
    def reset_preview():st.session_state.pop(editor_key,None)
    st.button('입력 되돌리기 · 저장된 값',key=f'reset_{editor_key}',on_click=reset_preview)
    if draft_grid!=original_grid:st.info('미저장 측정값이 있습니다. 아래에서 변경자를 입력하고 저장해주세요.')
    with st.form(f'preview_save_{kind}_{r["id"]}_{r["_version"]}'):
        actor=st.text_input('수치 변경자 *')
        if st.form_submit_button('편집한 측정값 저장'):
            try:
                item={**r}
                if spacing:item['values']=draft_values
                else:item['grid']=draft_grid
                store.save_record(path,kind,item,actor,r['_version'])
                st.session_state['hw_saved']='측정값을 저장했습니다.';st.rerun()
            except Exception as exc:st.error(f'저장 실패: {exc}')
    if spacing:table([{'차이':k,'측정값(mils)':v} for k,v in store.spacing_deltas(draft_values).items()])
    else:
        if r.get('maxMinSpan') is not None and r['maxMinSpan']!=store.flatness_span(r['grid']):st.info('백업에 저장된 Max−Min과 현재 측정값 계산이 다릅니다. 원본값을 자동 변경하지 않았습니다.')
        with st.expander('Flatness 3D · 낮은 값 파랑 / 높은 값 빨강'):
            surface_colors=[[0,'#64748b'],[1,'#64748b']] if stats['최소값']==stats['최댓값'] else [[0,MEASURE_BLUE],[1,MEASURE_RED]]
            surface=go.Figure(go.Surface(z=draft_grid,x=list(range(1,8)),y=list(range(1,6)),colorscale=surface_colors,connectgaps=False,lighting={'ambient':.65,'diffuse':.75,'specular':.18,'roughness':.85},lightposition={'x':100,'y':100,'z':150},colorbar={'title':'mils'}))
            surface.update_layout(height=420,scene={'xaxis_title':'Col · G/V → Window','yaxis_title':'Row','zaxis_title':'mils'})
            st.plotly_chart(surface,use_container_width=True)
    with st.expander('측정 정보 · 날짜/SN/비고 수정',expanded=st.session_state.get('measurement_edit_requested')==r['id']):
        measurement_form(path,kind,equipment,chambers,all_logs,r)
    with st.expander('선택 측정의 필드별 수정 기록 · 상세 로그'):show_audit(path,r['id'])
    measurement_export(rows,kind,spacing)


def export_buttons(df,name):
    a,b=st.columns(2)
    a.download_button('CSV 다운로드',safe_export(df).to_csv(index=False).encode('utf-8-sig'),name+'.csv','text/csv',key=name+'_csv')
    buf=BytesIO()
    with pd.ExcelWriter(buf,engine='openpyxl') as w:safe_export(df).to_excel(w,index=False,sheet_name='조회 결과')
    b.download_button('Excel 다운로드',buf.getvalue(),name+'.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key=name+'_xlsx')


def format_hw_workbook(content):
    """기존 내보내기도 같은 3×3/5×7 비교표로 변환합니다."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    book=load_workbook(BytesIO(content))
    if '챔버 정보' in book.sheetnames:del book['챔버 정보']
    for name,remove in [('장착 부품',{'관련 작업 ID','상태','계산 근거'}),('작업 교체 이력',{'원문 출처'}),('삭제 보관 작업',{'원문 출처'})]:
        if name not in book.sheetnames:continue
        ws=book[name]
        for col in range(ws.max_column,0,-1):
            if ws.cell(1,col).value in remove:ws.delete_cols(col)
        ws.auto_filter.ref=ws.dimensions
        for col in range(1,ws.max_column+1):
            title=ws.cell(1,col).value
            ws.column_dimensions[get_column_letter(col)].width=28 if title in ['S/N','교체 전 S/N','교체 후 S/N'] else 20
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    edge=Side(style='thin',color='94A3B8')
    def cell(ws,r,c,value,header=False):
        out=ws.cell(r,c,value)
        out.alignment=Alignment(horizontal='center' if header or isinstance(value,(int,float)) else 'left',vertical='center',wrap_text=True)
        out.border=Border(left=edge,right=edge,top=edge,bottom=edge)
        out.fill=PatternFill('solid',fgColor='CBD5E1' if header else 'F1F5F9')
        out.font=Font(name='맑은 고딕',size=11,bold=header,color='172B4D')
        if isinstance(value,str):out.data_type='s'
        return out
    for name in ['Spacing','Diffuser Flatness','Susceptor Flatness']:
        if name not in book.sheetnames:continue
        original=book[name];headers=[c.value for c in original[1]]
        records=[dict(zip(headers,row)) for row in original.iter_rows(min_row=2,values_only=True)]
        idx=book.sheetnames.index(name);del book[name];ws=book.create_sheet(name,idx)
        spacing=name=='Spacing';width=3 if spacing else 7;groups={}
        for r in records:groups.setdefault((r.get('설비'),r.get('챔버')),[]).append(r)
        ws.cell(1,1,name+' 실측 비교표 · 단위 mils · 0은 실측값 / 빈칸은 미측정')
        ws.cell(2,1,'같은 설비·챔버의 측정 기록을 오른쪽 열 방향으로 나란히 표시합니다.')
        ws.column_dimensions['A'].width=16
        base=4
        for (eq,ch),items in groups.items():
            ws.cell(base,1,f'{eq} / {ch}').font=Font(bold=True,size=13,color='173C72')
            for n,record in enumerate(items):
                start=2+n*(width+1)
                for c in range(start,start+width):ws.column_dimensions[get_column_letter(c)].width=15
                ws.column_dimensions[get_column_letter(start+width)].width=3
                metadata=['측정 ID','측정일','측정 구분','부품 S/N','측정기','비고']
                for offset,label in enumerate(metadata,1):
                    cell(ws,base+offset,1,label,True)
                    ws.merge_cells(start_row=base+offset,start_column=start,end_row=base+offset,end_column=start+width-1)
                    cell(ws,base+offset,start,record.get(label))
                    ws.row_dimensions[base+offset].height=42 if label=='비고' else 24
                top=base+8
                for ci in range(width):cell(ws,top,start+ci,(['좌','중앙','우'][ci] if spacing else f'Col {ci+1}'),True)
                matrix=store.POINT_LAYOUT if spacing else [[None]*7 for _ in range(5)]
                for ri,row in enumerate(matrix):
                    cell(ws,top+1+ri,1,(['G/V 측','중앙','Window 측'][ri] if spacing else f'Row {ri+1}'),True)
                    ws.row_dimensions[top+1+ri].height=34
                    for ci,point in enumerate(row):
                        key=f'P{point} {SPACING_LABELS[point]}' if spacing else f'R{ri+1}C{ci+1}'
                        cell(ws,top+1+ri,start+ci,record.get(key))
                summary=top+2+len(matrix)
                keys=list(store.spacing_deltas([None]*9)) if spacing else ['Max−Min','원본 Max−Min']
                for offset,key in enumerate(keys):
                    cell(ws,summary+offset,1,key,True)
                    ws.merge_cells(start_row=summary+offset,start_column=start,end_row=summary+offset,end_column=start+width-1)
                    cell(ws,summary+offset,start,record.get(key))
            base+=24
        if not records:ws.cell(4,1,'등록된 측정 데이터 없음')
        ws.freeze_panes='B4';ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=0;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:2'
    if '내보내기 안내' in book.sheetnames:
        ws=book['내보내기 안내'];ws.append(['측정표 배치','같은 설비·챔버별 기록을 열 방향으로 비교. Spacing 3×3 / Flatness 5×7.'])
    out=BytesIO();book.save(out);return out.getvalue()


def hw_excel_bytes(path):
    """H/W 전용 조회용 Excel. 같은 읽기 트랜잭션의 레코드와 변경 이력을 사용합니다."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    conn=store.connection(path)
    try:
        conn.execute('BEGIN')
        snapshot={kind:[] for kind in store.KINDS}
        for row in conn.execute('SELECT kind,payload FROM entities ORDER BY rowid'):
            snapshot[row['kind']].append(json.loads(row['payload']))
        audit=[dict(row) for row in conn.execute('SELECT * FROM audit ORDER BY timestamp,id')]
    finally:conn.close()
    sheets={}
    def add(name,rows,columns):sheets[name]=pd.DataFrame(rows,columns=columns)
    def mapped(item,mapping):return {label:item.get(field,'') for field,label in mapping.items()}
    def serialized(value):return json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else value
    eqs=snapshot['equipments'];logs=[r for r in snapshot['workLogs'] if not r.get('deletedAt')]
    created=datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M:%S KST')
    add('내보내기 안내',[['생성 시각',created],['범위','현재 화면 필터와 무관하게 저장된 전체 설비 H/W 데이터'],['단위','측정값은 mils. 0은 실측값, 빈칸은 미측정'],['보관 기록','삭제 보관 작업은 별도 시트. 장착 부품은 활성 작업 이력 계산 기준'],['첨부','첨부파일 목록만 포함하며 파일 본문은 JSON 또는 통합 ZIP 백업에 포함'],['용도','조회용 Excel이며 이 파일로 DB 복원하거나 자동 재등록하지 않습니다. Pending·REV 제외']],['항목','설명'])
    add('설비 정보',[mapped(r,{'id':'설비 ID','name':'설비명','customer':'고객사'}) for r in eqs],['설비 ID','설비명','고객사'])
    chambers=[];installed=[]
    for eq in eqs:
        for ch in eq['chambers']:
            chambers.append({'설비':eq['id'],'챔버':ch['chamberId'],'챔버명':ch.get('name',''),'초기 장착 원문':serialized(ch.get('initialParts',{}))})
            for part in store.installed_parts(eq,ch['chamberId'],logs):
                installed.append({'설비':eq['id'],'챔버':ch['chamberId'],**mapped(part,{'partType':'부품','sn':'S/N','installedDate':'장착일','lastWorkDate':'최근 작업일','lastWorkLogId':'관련 작업 ID','status':'상태','basis':'계산 근거'})})
    add('챔버 정보',chambers,['설비','챔버','챔버명','초기 장착 원문'])
    add('장착 부품',installed,['설비','챔버','부품','S/N','장착일','최근 작업일','관련 작업 ID','상태','계산 근거'])
    work_map={'id':'작업 ID','equipmentId':'설비','chamberId':'챔버','workDate':'작업일','workTimestamp':'작업 시각','workOrder':'작업 차수','workType':'작업 유형','targetPart':'부품','beforeSn':'교체 전 S/N','afterSn':'교체 후 S/N','reason':'사유','details':'내용','worker':'작업자','issueStatus':'결과','isAmbiguous':'원본 검토 표시','ambiguityReason':'원본 검토 사유','pendingLink':'Pending 연결','source':'원문 출처'}
    add('작업 교체 이력',[mapped(r,work_map) for r in logs],list(work_map.values()))
    trash_map={**work_map,'deletedAt':'삭제 시각','deletedBy':'삭제자','deletionReason':'삭제 사유'}
    add('삭제 보관 작업',[mapped(r,trash_map) for r in snapshot['workLogs'] if r.get('deletedAt')],list(trash_map.values()))
    base={'id':'측정 ID','equipmentId':'설비','chamberId':'챔버','measuredAt':'측정일','timing':'측정 구분','gaugeTool':'측정기','unit':'단위','relatedWorkLogId':'관련 작업 ID','notes':'비고','isAmbiguous':'원본 검토 표시','source':'원문 출처'}
    def measurement(r):return {**mapped(r,base),'부품 S/N':r.get('partSnAtMeasurement') or r.get('partSn') or r.get('diffuserSn') or r.get('susceptorSn') or ''}
    spacing=[]
    for r in snapshot['spacingMeasurements']:
        values=r['values'];spacing.append({**measurement(r),**{f'P{i} {SPACING_LABELS[i]}':v for i,v in enumerate(values,1)},**store.spacing_deltas(values)})
    add('Spacing',spacing,list(base.values())+['부품 S/N']+[f'P{i} {SPACING_LABELS[i]}' for i in range(1,10)]+list(store.spacing_deltas([None]*9)))
    for kind in ['Diffuser','Susceptor']:
        rows=[]
        for r in snapshot['flatnessMeasurements']:
            if r['type']!=kind:continue
            rows.append({**measurement(r),'Max−Min':store.flatness_span(r['grid']),'원본 Max−Min':r.get('maxMinSpan'),**{f'R{ri}C{ci}':v for ri,row in enumerate(r['grid'],1) for ci,v in enumerate(row,1)}})
        add(kind+' Flatness',rows,list(base.values())+['부품 S/N','Max−Min','원본 Max−Min']+[f'R{ri}C{ci}' for ri in range(1,6) for ci in range(1,8)])
    add('측정 위치 기준',[mapped(r,{'id':'ID','pointIndex':'포인트','label':'위치','description':'설명'}) for r in snapshot['spacingStandards']],['ID','포인트','위치','설명'])
    add('변경 이력',[mapped(r,{'timestamp':'시각(UTC)','actor':'변경자','action':'작업','kind':'종류','entity_id':'항목 ID','before_json':'변경 전','after_json':'변경 후'}) for r in audit],['시각(UTC)','변경자','작업','종류','항목 ID','변경 전','변경 후'])
    files=[]
    for r in snapshot['workLogs']:
        for a in r.get('attachments') or []:
            if isinstance(a,dict):files.append({'작업 ID':r['id'],'설비':r['equipmentId'],'챔버':r['chamberId'],'파일명':a.get('name',''),'유형':a.get('type',''),'크기(bytes)':a.get('size',''),'파일 ID':a.get('blobId',''),'원본 경로':a.get('url','')})
    add('첨부파일 목록',files,['작업 ID','설비','챔버','파일명','유형','크기(bytes)','파일 ID','원본 경로'])
    buf=BytesIO()
    with pd.ExcelWriter(buf,engine='openpyxl') as writer:
        for name,frame in sheets.items():
            for col in frame.columns:
                frame[col]=frame[col].map(lambda v:ILLEGAL_CHARACTERS_RE.sub('',str(serialized(v))) if isinstance(v,(str,list,dict)) else v)
            safe_export(frame).to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.fill=PatternFill('solid',fgColor='CBD5E1');cell.font=Font(bold=True,color='172B4D')
            for row in ws:
                for cell in row:
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    cell.border=Border(bottom=Side(style='thin',color='94A3B8'))
            for column in ws.columns:
                letter=column[0].column_letter
                ws.column_dimensions[letter].width=min(48,max(14,max(len(str(c.value or '')) for c in list(column)[:100])+3))
    return format_hw_workbook(buf.getvalue())


def full_backup(data_dir):
    buf=BytesIO()
    with tempfile.TemporaryDirectory(dir=data_dir) as tmp:
        with ZipFile(buf,'w',ZIP_DEFLATED) as z:
            for name in ['pending.sqlite3','hw_history.sqlite3']:
                source=data_dir/name
                if source.exists():
                    target=Path(tmp)/name
                    conn=sqlite3.connect(f'{source.resolve().as_uri()}?mode=ro',uri=True)
                    dest=sqlite3.connect(target)
                    try:conn.backup(dest)
                    finally:conn.close();dest.close()
                    z.write(target,'data/'+name)
            for name in ['default.xlsx','source.xlsx']:
                if (data_dir/name).exists():z.write(data_dir/name,'data/'+name)
            z.writestr('RESTORE.txt','앱 종료 후 data 폴더를 먼저 백업하고 이 ZIP의 data 폴더를 복사하세요. 코드와 가상환경은 별도 설치가 필요합니다. DB와 Excel은 같은 백업 묶음을 사용하세요.')
    return buf.getvalue()


def data_management(path,data_dir):
    st.subheader('데이터 가져오기 · 백업')
    st.caption('JSON은 가져오기 전에 검증합니다. 동일 ID가 이미 있으면 기존 기록을 보존하고 건너뜁니다. 이 화면은 외부 서버에 접속하지 않습니다.')
    upload=st.file_uploader('H/W JSON 백업 업로드',type=['json'],key='hw_import')
    if upload:
        try:
            current={k:store.records(path,k,include_deleted=True) for k in store.KINDS}
            plan=store.parse_import(upload.getvalue(),current)
            table([{'종류':store.KINDS[k],'가져올 유효 항목':len(rows),'기존 ID 중복':sum(r['id'] in {e['id'] for e in current[k]} for r in rows)} for k,rows in plan['records'].items()])
            if plan['mode']=='sample':st.warning('백업의 모드가 sample입니다. 실제 원본과 대조한 데이터인지 검토해주세요. 자동으로 실제 데이터라고 확정하지 않습니다.')
            if not plan['native']:st.info('AI Studio 백업의 독립된 과거 수정·삭제 이력 목록은 포함되어 있지 않습니다. 각 항목의 원문 audit 필드는 보존됩니다.')
            if plan['errors']:
                st.error(f'가져올 수 없는 항목 {len(plan["errors"])}건');table(plan['errors'])
            if plan['warnings']:
                with st.expander(f'원본 검토 및 첨부 경로 안내 {len(plan["warnings"])}건'):table(plan['warnings'])
            actor=st.text_input('가져오기 담당자',key='hw_import_actor')
            approved=st.checkbox('데이터 출처와 미리보기를 검토했습니다.',key='hw_import_reviewed')
            partial=st.checkbox('오류 항목은 제외하고 유효 항목만 가져옵니다.',key='hw_import_partial') if plan['errors'] else False
            if st.button('검토한 데이터 가져오기',disabled=not approved):
                result=store.import_backup(path,upload.getvalue(),actor,partial)
                st.success(f"새로 가져옴 {result['added']}건 / 기존 ID 보존 {result['skipped']}건 / 오류 제외 {result['errors']}건")
                st.caption('위쪽 메뉴에서 설비 현황이나 작업 이력으로 이동해 확인하세요.')
        except Exception as exc:st.error(f'가져오기 실패: {exc}')
    st.divider()
    st.download_button('H/W 전체 백업 JSON · 첨부·변경 이력 포함',store.backup_bytes(path),'HW_Full_Backup_'+today().isoformat()+'.json','application/json')
    st.caption('H/W 백업 JSON을 위 입력창에 올리면 비어 있는 저장소로 복원할 수 있습니다. 기존 ID는 덮어쓰지 않습니다.')
    if st.button('Pending·REV·H/W 통합 백업 ZIP 준비'):
        st.session_state['hw_full_backup']=full_backup(data_dir)
        st.session_state['hw_full_backup_time']=datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M:%S KST')
    if st.session_state.get('hw_full_backup'):
        st.caption('백업 생성 시각: '+st.session_state['hw_full_backup_time']+' · 이후 변경사항은 준비 버튼을 다시 눌러 반영하세요.')
        st.download_button('통합 백업 ZIP 다운로드',st.session_state['hw_full_backup'],'Pending_HW_Backup_'+today().isoformat()+'.zip','application/zip')
    st.caption('통합 ZIP 복원: 앱을 종료하고 기존 data 폴더를 백업한 뒤 ZIP 안의 data 폴더를 복사합니다. 서로 다른 시각에 만든 DB와 Excel을 섞지 마세요. Temp의 파일은 정리될 수 있으므로 회사에서 허용한 위치에 백업하세요.')
    with st.expander('H/W 전체 변경 이력'):show_audit(path)


def render_hw(data_dir,context=None):
    path=data_dir/'hw_history.sqlite3'
    st.markdown("""<style>
.hw-parts-scroll {overflow:auto;border:1px solid #94a3b8;border-radius:8px;margin-bottom:16px}
.hw-parts-table {width:100%;border-collapse:collapse;background:#f1f5f9;min-width:620px;table-layout:fixed;font-size:13px}
.hw-parts-table th {background:#cbd5e1;text-align:left;padding:10px;border:1px solid #94a3b8}
.hw-parts-table th:nth-child(1) {width:19%}.hw-parts-table th:nth-child(2) {width:27%}.hw-parts-table th:nth-child(3),.hw-parts-table th:nth-child(4) {width:20%}.hw-parts-table th:nth-child(5) {width:14%}
.hw-parts-table td {overflow-wrap:anywhere;padding:10px;border:1px solid #b6c4d5}
.hw-work-badges {display:flex;align-items:center;flex-wrap:wrap;gap:5px;min-height:30px}
.hw-work-badges span {background:#dbeafe;border:1px solid #94a3b8;color:#173c72;padding:2px 6px;border-radius:4px;font-size:11px;font-weight:600}
.hw-work-badges .hw-work-date {background:transparent;border:0;padding-left:0;color:#075985}
.hw-work-badges .hw-work-result {background:#d1fae5;border-color:#6da997;color:#065f46}
.hw-work-sn {background:#e2e8f0;border:1px solid #94a3b8;border-radius:5px;padding:7px 10px;color:#173c72;font-size:12px;line-height:1.5;overflow-wrap:anywhere}
.hw-work-body {border-top:1px solid #b6c4d5;margin-top:12px;padding-top:12px;color:#172b4d;font-size:13px;line-height:1.6}
.hw-work-reason,.hw-work-description {white-space:pre-wrap;overflow-wrap:anywhere;margin:0 0 7px}
.hw-work-worker {border-top:1px solid #cbd5e1;margin-top:10px;padding:9px 0 12px;color:#475569;font-size:12px}
[class*="st-key-hw_work_card_"] {background:#f1f5f9;border-radius:10px}
[class*="st-key-hw_work_card_"] [data-testid="stVerticalBlock"] {gap:9px}
[class*="st-key-hw_work_card_"] [data-testid="stHorizontalBlock"] {gap:8px;align-items:center}
[class*="st-key-hw_work_card_"] button {min-height:30px;padding:3px 8px}
[class*="st-key-hw_work_card_"] button p {font-size:12px}
</style>""",unsafe_allow_html=True)
    st.header('설비 H/W 이력관리')
    st.caption('로컬 저장 · 기존 Pending·REV 저장소 보존 · Firebase 자동 접속 없음')
    with st.expander('설비 H/W 전체 Excel 내보내기'):
        st.caption('저장된 전체 H/W 데이터를 항목별 시트로 내보냅니다. Pending·REV와 첨부파일 본문은 제외합니다.')
        if st.button('H/W Excel 파일 준비',key='hw_excel_prepare'):
            try:
                st.session_state['hw_excel_download']=hw_excel_bytes(path)
                st.session_state['hw_excel_created']=datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M:%S KST')
            except Exception as exc:st.error(f'Excel 생성 실패: {exc}')
        if st.session_state.get('hw_excel_download'):
            st.caption('생성 시각: '+st.session_state['hw_excel_created']+' · 이후 수정한 내용은 준비 버튼을 다시 눌러 반영하세요.')
            st.download_button('H/W Excel 다운로드',st.session_state['hw_excel_download'],'HW_History_'+today().isoformat()+'.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='hw_excel_download_button')

    conn=store.connection(path)
    try:sample_import=bool(conn.execute("SELECT 1 FROM imports WHERE mode='sample'").fetchone())
    finally:conn.close()
    if sample_import:st.warning('가져온 데이터에 sample 모드로 내보낸 백업이 포함되어 있습니다. 실제 원본과 대조 전에는 검증된 운영 데이터로 간주하지 마세요.')
    saved_message=st.session_state.pop('hw_saved',None)
    if saved_message:
        st.success(saved_message if isinstance(saved_message,str) else '저장했습니다. 변경 이력을 기록했습니다.')
        st.session_state.pop('hw_edit_id',None)
    section=card_menu('H/W 메뉴',['설비 현황','작업 이력','부품 추적','Spacing','Flatness','데이터 관리'],key='hw_section',horizontal=True)
    if section=='데이터 관리':data_management(path,data_dir);return
    equipments=store.records(path,'equipments')
    with st.expander('＋ 설비 등록'):
        with st.form('new_equipment'):
            identifier=st.text_input('설비 ID * · 예: A4CVD15I')
            name=st.text_input('설비 이름 *')
            customer=st.text_input('고객사')
            chamber_names=st.text_input('챔버 ID · 쉼표로 구분',value='Chamber A, Chamber B, Chamber C, Chamber D, Chamber E')
            actor=st.text_input('등록자 *')
            if st.form_submit_button('설비 저장'):
                try:
                    if not name.strip():raise ValueError('설비 이름을 입력해주세요.')
                    item={'id':identifier.strip(),'name':name.strip(),'customer':customer,'chambers':[{'chamberId':c.strip(),'name':c.strip(),'initialParts':{},'installedParts':{}} for c in chamber_names.split(',') if c.strip()]}
                    store.save_record(path,'equipments',item,actor)
                    st.session_state['hw_saved']=True;st.rerun()
                except Exception as exc:st.error(str(exc))
    if not equipments:
        st.info('아직 H/W 데이터가 없습니다. 「데이터 관리」에서 받은 JSON을 검토 후 가져오거나 설비부터 등록해주세요.');return
    lookup={r['id']:r for r in equipments}
    prefill=st.session_state.get('pending_hw_prefill',{})
    preferred=prefill.get('equipmentId')
    if preferred in lookup and 'hw_equipment' not in st.session_state:st.session_state['hw_equipment']=preferred
    equipment=st.selectbox('H/W 설비',list(lookup),format_func=lambda x:f'{x} · {lookup[x].get("name","")}',key='hw_equipment')
    chambers=[c['chamberId'] for c in lookup[equipment]['chambers']]
    chamber=st.selectbox('H/W 챔버',['전체']+chambers,key=f'hw_chamber_{equipment}')
    logs=store.records(path,'workLogs')
    if section=='설비 현황':
        for ch in chambers if chamber=='전체' else [chamber]:
            st.subheader(ch)
            parts=store.installed_parts(lookup[equipment],ch,logs)
            installed_table(parts)
        st.caption('초기 장착 기준에 시간순 부품 교체 이력을 적용한 결과입니다. 검토 표시가 있는 작업은 자동 반영하지 않습니다. 기록 불일치는 원본과 대조해주세요.')
        st.caption('색상 기준: 장착일·최근 변경일 각각 0~364일은 초록→노랑→주황, 365일 이상은 빨강입니다. 경과 표시는 최근 변경일(없으면 장착일) 기준입니다. 원본처럼 RFG·Matcher·RPSC Reactor·MF Generator는 색상에서 제외합니다.')
        monthly_chart(logs,equipment,chambers)
        with st.expander('설비 이름·고객사 수정'):
            eq=lookup[equipment]
            with st.form(f'equipment_edit_{equipment}_{eq["_version"]}'):
                new_name=st.text_input('설비 이름',value=eq.get('name',''))
                new_customer=st.text_input('고객사',value=eq.get('customer',''))
                operator=st.text_input('변경자 *')
                if st.form_submit_button('설비 정보 수정 저장'):
                    try:
                        store.save_record(path,'equipments',{**eq,'name':new_name,'customer':new_customer},operator,eq['_version'])
                        st.session_state['hw_saved']=True;st.rerun()
                    except Exception as exc:st.error(str(exc))
    elif section=='작업 이력':
        query=st.text_input('작업 이력 검색 · 부품/SN/내용/담당자',key='hw_work_search')
        with st.expander('＋ 작업 이력 등록',expanded=bool(st.session_state.get('pending_hw_prefill'))):work_form(path,equipment,chambers,context)
        rows=filter_rows(logs,equipment,chamber,query)
        work_trash(path,equipment)
        if not rows:st.info('해당 조건의 작업 이력이 없습니다.');return
        df=pd.DataFrame([{'ID':r['id'],'설비':r['equipmentId'],'챔버':r['chamberId'],'작업일':r['workDate'],'작업':r['workType'],'부품':r.get('targetPart',''),'교체 전 S/N':r.get('beforeSn',''),'교체 후 S/N':r.get('afterSn',''),'내용':r.get('details',''),'작업자':r.get('worker',''),'결과':r.get('issueStatus','')} for r in rows])
        work_cards(path,rows,equipment,chambers,context)
        with st.expander('표로 보기 · Excel/CSV 다운로드'):
            st.dataframe(df,hide_index=True,use_container_width=True);export_buttons(df,'hw_work_logs')
        indexed={r['id']:r for r in rows}
        if st.session_state.get('hw_work_detail') not in indexed:st.session_state['hw_work_detail']=next(iter(indexed))
        selected=st.selectbox('작업 상세 선택',list(indexed),key='hw_work_detail',format_func=lambda x:f"{indexed[x]['workDate']} · {indexed[x].get('targetPart','')} · {x}")
        item=indexed[selected]
        st.write(item.get('details',''));st.caption('작업 사유: '+item.get('reason',''))
        if item.get('isAmbiguous'):st.warning(item.get('ambiguityReason') or '원본 검토 표시가 있는 기록입니다.')
        if item.get('source'):st.json(item['source'])
        show_pending_link(item.get('pendingLink'),context)
        attachments(path,item)
        if st.session_state.get('hw_edit_id')!=item['id']:
            with st.expander('작업 이력 수정'):work_form(path,equipment,chambers,context,item)
        with st.expander('이 작업의 변경 이력'):show_audit(path,item['id'])
    elif section=='부품 추적':
        st.caption('부품 이동 확인을 위해 전체 설비의 동일 S/N 기록을 조회합니다.')
        sn=st.text_input('조회할 부품 S/N · 정확히 일치하는 번호',key='hw_trace_sn')
        if sn:
            traced=store.trace_sn(equipments,logs,sn.strip())
            table(traced)
            if traced:export_buttons(pd.DataFrame(traced),'hw_part_trace')
        else:st.info('부품 S/N을 입력하면 설비·챔버 간 기록을 확인할 수 있습니다.')
    else:measurement_view(path,'spacingMeasurements' if section=='Spacing' else 'flatnessMeasurements',equipment,chamber,chambers,logs)


def installed_table(parts):
    from html import escape
    rows=[]
    headers=['부품','S/N','장착일','최근 변경일','경과']
    for p in parts:
        recent=p.get('lastWorkDate') or p.get('installedDate','')
        age=store.installation_age(recent,p['partType'],today())
        installed_age=store.installation_age(p.get('installedDate',''),p['partType'],today())
        cells=[p['partType'],p.get('sn',''),p.get('installedDate',''),recent,age['label']]
        rows.append('<tr>'+''.join(f'<td style="{installed_age["style"] if i==2 else age["style"] if i in [3,4] else ""}">'+escape(str(v))+'</td>' for i,v in enumerate(cells))+'</tr>')
    st.markdown('<div class="hw-parts-scroll"><table class="hw-parts-table"><thead><tr>'+''.join('<th>'+h+'</th>' for h in headers)+'</tr></thead><tbody>'+''.join(rows)+'</tbody></table></div>',unsafe_allow_html=True)


def monthly_chart(logs,equipment,chambers):
    st.subheader('월별 작업 · 부품 교체 건수')
    eq_logs=[r for r in logs if r['equipmentId']==equipment]
    years=sorted({int(r['workDate'][:4]) for r in eq_logs if store.good_date(r.get('workDate'))}|{today().year},reverse=True)
    with st.container(border=True):
        a,b,c=st.columns(3)
        replacement_years=[int(r['workDate'][:4]) for r in eq_logs if store.good_date(r.get('workDate')) and store.is_replacement(r)]
        existing_years=[int(r['workDate'][:4]) for r in eq_logs if store.good_date(r.get('workDate'))]
        default_year=max(replacement_years or existing_years or [today().year])
        year=a.selectbox('집계 연도',years,index=years.index(default_year),key=f'monthly_year_{equipment}')
        chamber=b.selectbox('집계 Chamber',['전체 Chamber']+chambers,key=f'monthly_chamber_{equipment}')
        part=c.selectbox('교체 부품 · 표준 8종',['전체 부품']+store.PARTS,key=f'monthly_part_{equipment}')
        a,b=st.columns(2)
        mode=a.selectbox('집계 대상',['부품 교체','전체 작업'],key=f'monthly_mode_{equipment}')
        status=b.selectbox('집계 작업 결과',['전체','정상완료','조치필요','모니터링중','이슈발생'],key=f'monthly_status_{equipment}')
        counts,selected=store.monthly_work(logs,equipment,year,chamber,part,mode,status)
        frame=pd.DataFrame(counts).set_index('월')
        st.caption(f'{year}년 / {chamber} / {part} · {len(selected)}건 · 월별 필터는 위 장착 현황의 챔버 선택과 별도로 적용됩니다.')
        st.bar_chart(frame[store.PARTS if part=='전체 부품' else [part]])
        st.caption('초기 셋업·점검은 부품 교체 집계에서 제외합니다. 전체 작업을 선택하면 포함합니다. 같은 S/N의 PM도 교체/PM 유형이면 1건으로 집계합니다.')
        with st.expander('월별 집계 표와 대상 이력'):
            st.dataframe(frame,use_container_width=True)
            table([{'ID':r['id'],'날짜':r['workDate'],'챔버':r['chamberId'],'부품':r.get('targetPart',''),'유형':r.get('workType',''),'결과':r.get('issueStatus','')} for r in selected])
            export_buttons(frame.reset_index(),'hw_monthly')


def choose_work(identifier,edit=False):
    st.session_state['hw_work_detail']=identifier
    st.session_state['hw_edit_id']=identifier if edit else None


def trace_work(sn):
    st.session_state['hw_trace_sn']=sn
    st.session_state['hw_section']='부품 추적'


def request_archive(identifier):
    st.session_state['hw_archive_id']=identifier


def work_cards(path,rows,equipment,chambers,context):
    from html import escape
    rows=sorted(rows,key=lambda r:(r.get('workTimestamp') or r['workDate'],r.get('workOrder') or 1,r['id']),reverse=True)
    pages=max(1,math.ceil(len(rows)/10))
    page=st.number_input('작업 카드 페이지 · 10건씩',min_value=1,max_value=pages,value=1,step=1,key=f'work_page_{equipment}_{len(rows)}')
    st.caption(f'작업 및 교체 이력 {len(rows)}건 · 카드의 기능 버튼을 눌러 실행하세요.')
    for item in rows[(page-1)*10:page*10]:
        identifier=item['id']
        with st.container(border=True,key=f'hw_work_card_{identifier}'):
            header,result,edit,delete=st.columns([8,1.4,1,1],gap='small')
            badges=[item.get('workTimestamp') or item['workDate'],f"{item.get('workOrder') or 1}차 작업",item['chamberId'],item.get('workType',''),item.get('targetPart','')]
            badge_html=''.join(f'<span class="{"hw-work-date" if i==0 else "hw-work-result" if i==5 and x=="정상완료" else ""}">{escape(str(x))}</span>' for i,x in enumerate(badges) if x)
            header.markdown('<div class="hw-work-badges">'+badge_html+'</div>',unsafe_allow_html=True)
            result.markdown('<div class="hw-work-badges"><span class="'+('hw-work-result' if item.get('issueStatus')=='정상완료' else '')+'">'+escape(str(item.get('issueStatus') or '미기재'))+'</span></div>',unsafe_allow_html=True)
            edit.button('수정',key=f'card_edit_{identifier}',on_click=choose_work,args=(identifier,True),use_container_width=True)
            delete.button('삭제',key=f'card_archive_{identifier}',on_click=request_archive,args=(identifier,),use_container_width=True)
            serial,trace=st.columns([8,2],gap='small')
            serial.markdown('<div class="hw-work-sn">교체 전 S/N: <b>'+escape(str(item.get('beforeSn') or '미기재'))+'</b> → 교체 후 S/N: <b>'+escape(str(item.get('afterSn') or '미기재'))+'</b></div>',unsafe_allow_html=True)
            sn=item.get('afterSn') or item.get('beforeSn') or ''
            trace.button('부품 이동 추적',key=f'card_trace_{identifier}',on_click=trace_work,args=(sn,),disabled=not sn,use_container_width=True)
            body='<div class="hw-work-body"><div class="hw-work-reason"><b>사유:</b> '+escape(str(item.get('reason') or '미기재'))+'</div>'
            if item.get('details'):body+='<div class="hw-work-description"><b>내용:</b> '+escape(str(item['details']))+'</div>'
            body+='<div class="hw-work-worker">작업자: <b>'+escape(str(item.get('worker') or '미기재'))+'</b></div></div>'
            st.markdown(body,unsafe_allow_html=True)
            if st.session_state.get('hw_edit_id')==identifier:
                st.subheader('선택한 작업 수정')
                work_form(path,equipment,chambers,context,item)
            if st.session_state.get('hw_archive_id')==identifier:
                st.warning('삭제하면 장착 현황 계산·월별 집계·기본 목록에서 제외됩니다. 원문과 첨부·측정 기록은 보존하며 삭제 보관함에서 복원할 수 있습니다.')
                with st.form(f'archive_form_{identifier}_{item["_version"]}'):
                    actor=st.text_input('삭제 담당자 *')
                    reason=st.text_input('삭제 사유 *')
                    if st.form_submit_button('삭제 보관 확인'):
                        try:
                            store.archive_work(path,identifier,actor,reason,item['_version'])
                            st.session_state.pop('hw_archive_id',None)
                            st.session_state.pop('hw_edit_id',None)
                            st.session_state['hw_saved']='작업을 삭제 보관했습니다.';st.rerun()
                        except Exception as exc:st.error(str(exc))
                st.button('삭제 취소',key=f'cancel_archive_{identifier}',on_click=lambda:st.session_state.pop('hw_archive_id',None))


def work_trash(path,equipment):
    rows=[r for r in store.records(path,'workLogs',include_deleted=True) if r.get('deletedAt') and r['equipmentId']==equipment]
    with st.expander(f'삭제 보관함 · {len(rows)}건'):
        if not rows:st.caption('보관된 작업이 없습니다.');return
        indexed={r['id']:r for r in rows}
        selected=st.selectbox('복원할 작업',list(indexed),format_func=lambda x:f"{indexed[x]['workDate']} · {indexed[x].get('targetPart','')} · {x}")
        item=indexed[selected]
        st.text('삭제 사유: '+item.get('deletionReason',''))
        with st.form(f'restore_work_{selected}_{item["_version"]}'):
            actor=st.text_input('복원 담당자 *')
            reason=st.text_input('복원 사유 *')
            if st.form_submit_button('작업 이력 복원'):
                try:
                    store.archive_work(path,selected,actor,reason,item['_version'],restore=True)
                    st.session_state['hw_saved']='작업을 복원했습니다.';st.rerun()
                except Exception as exc:st.error(str(exc))


MEASURE_BLUE='#2563eb'
MEASURE_RED='#ef4444'
SPACING_LABELS={1:'G/V 좌',6:'G/V 중앙',2:'G/V 우',9:'중앙 좌',5:'Center 기준',7:'중앙 우',4:'Window 좌',8:'Window 중앙',3:'Window 우'}


def measurement_number(value):
    return '미측정' if value is None else f'{value:g}'


def flatness_color(value,low,high):
    if value is None:return '#e2e8f0'
    if low==high:return '#64748b'
    ratio=max(0,min(1,(value-low)/(high-low)))
    a=(37,99,235);b=(239,68,68)
    return '#'+''.join(f'{round(x+(y-x)*ratio):02x}' for x,y in zip(a,b))


def flatness_table_html(grid):
    from html import escape
    labels=['G/V 상단','상단','Center','하단','Window 하단']
    parts=['<table class="flatness-values" aria-label="Flatness 실측 데이터 표"><thead><tr><th>구분</th>']
    parts.extend(f'<th>Col {col}</th>' for col in range(1,8))
    parts.append('</tr></thead><tbody>')
    for ri,row in enumerate(grid,1):
        parts.append(f'<tr><th>Row {ri}<small>{labels[ri-1]}</small></th>')
        for ci,value in enumerate(row,1):
            parts.append(f'<td><small>R{ri}C{ci}</small><strong>{escape(measurement_number(value))}</strong></td>')
        parts.append('</tr>')
    parts.append('</tbody></table>')
    return ''.join(parts)


def spacing_table_html(values):
    from html import escape
    cells=[]
    for row in store.POINT_LAYOUT:
        for point in row:
            v=values[point-1];delta=None if v is None else round(v-885,3)
            color='#64748b' if v is None else '#1d4ed8' if delta<0 else '#b91c1c' if delta>0 else '#047857'
            cells.append(f'<div class="spacing-cell {"spacing-center" if point==5 else ""}"><small>P{point} · {SPACING_LABELS[point]}</small><strong>{escape(measurement_number(v))}</strong><span>mils</span><small style="color:{color}">885 대비 {"—" if delta is None else f"{delta:+g}"}</small></div>')
    return '<div class="spacing-values">'+''.join(cells)+'</div>'


def spacing_diagram_svg(values,table_mode=False):
    from html import escape
    svg=['<svg xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Spacing 9포인트 위치도" viewBox="0 0 700 510" class="measurement-svg">',
         '<rect x="35" y="40" width="630" height="420" rx="24" fill="#e5ecf5" stroke="#64748b" stroke-width="3"/>',
         '<rect x="62" y="67" width="576" height="366" rx="14" fill="#f1f5f9" stroke="#94a3b8" stroke-width="2"/>',
         '<path d="M350 68V432M64 250H636" stroke="#94a3b8" stroke-dasharray="4 5"/>',
         '<text x="350" y="24" text-anchor="middle" fill="#075985" font-size="18" font-weight="700">▲ GATE VALVE · G/V 측</text>',
         '<text x="350" y="490" text-anchor="middle" fill="#92400e" font-size="18" font-weight="700">▼ WINDOW 측</text>']
    if table_mode:
        svg=svg[:1]+['<rect x="62" y="67" width="576" height="366" rx="12" fill="#f1f5f9" stroke="#94a3b8"/>','<text x="350" y="24" text-anchor="middle" fill="#075985" font-size="18" font-weight="700">▲ GATE VALVE · G/V 측</text>','<text x="350" y="490" text-anchor="middle" fill="#92400e" font-size="18" font-weight="700">▼ WINDOW 측</text>']
    for y,row in enumerate(store.POINT_LAYOUT):
        for x,point in enumerate(row):
            left=82+x*185;top=90+y*112;v=values[point-1]
            stroke='#0891b2' if point==5 else '#94a3b8'
            svg.append(f'<rect x="{left}" y="{top}" width="166" height="86" rx="9" fill="{"#e0f2fe" if point==5 else "#ffffff"}" stroke="{stroke}" stroke-width="{3 if point==5 else 1.5}"/>')
            svg.append(f'<text x="{left+83}" y="{top+21}" text-anchor="middle" font-size="13" fill="#334155">P{point} · {SPACING_LABELS[point]}</text>')
            svg.append(f'<text x="{left+83}" y="{top+49}" text-anchor="middle" font-size="25" font-weight="700" fill="#172b4d">{escape(measurement_number(v))}</text>')
            delta='—' if v is None else f'{v-885:+.1f} mils'
            color='#64748b' if v is None else '#1d4ed8' if v<885 else '#b91c1c' if v>885 else '#047857'
            svg.append(f'<text x="{left+83}" y="{top+71}" text-anchor="middle" font-size="13" fill="{color}">{delta}</text>')
    return ''.join(svg)+'</svg>'


def flatness_diagram_svg(grid,color_range=None,table_mode=False):
    from html import escape
    vals=[v for row in grid for v in row if v is not None];low=min(vals) if vals else None;high=max(vals) if vals else None
    if color_range is not None:low,high=color_range
    out=['<svg xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Flatness 35포인트 위치도" viewBox="0 0 700 490" class="measurement-svg">',
         '<defs><linearGradient id="flatshade" x1="0" y1="0" x2="0.7" y2="1"><stop offset="0" stop-color="#ffffff" stop-opacity=".20"/><stop offset=".45" stop-color="#ffffff" stop-opacity="0"/><stop offset="1" stop-color="#172b4d" stop-opacity=".20"/></linearGradient></defs>',
         '<rect x="35" y="45" width="630" height="395" rx="22" fill="#d9e2ef" stroke="#475569" stroke-width="3"/>',
         '<rect x="53" y="62" width="594" height="360" rx="12" fill="#f1f5f9" stroke="#94a3b8" stroke-width="2"/>',
         '<text x="19" y="245" transform="rotate(-90 19 245)" text-anchor="middle" fill="#075985" font-size="16">Gate Valve 측 · Col 1</text>',
         '<text x="684" y="245" transform="rotate(90 684 245)" text-anchor="middle" fill="#92400e" font-size="16">Window 측 · Col 7</text>']
    for ri,row in enumerate(grid):
        for ci,v in enumerate(row):
            x=68+ci*80;y=82+ri*64
            color=('#ffffff' if v is not None else '#e2e8f0') if table_mode else flatness_color(v,low,high)
            out.append(f'<rect x="{x}" y="{y}" width="80" height="64" fill="{color}" stroke="#cbd5e1" stroke-width="1"/>')
            if not table_mode:out.append(f'<rect x="{x}" y="{y}" width="80" height="64" fill="url(#flatshade)"/>')
            if v is None:
                out.append(f'<path d="M{x+8} {y+54}L{x+68} {y+9}" stroke="#94a3b8" stroke-width="1.3"/>')
            out.append(f'<text x="{x+5}" y="{y+13}" font-size="9" fill="{"#475569" if v is None or table_mode else "#ffffff"}">R{ri+1}C{ci+1}</text>')
            label='—' if v is None else measurement_number(v)
            out.append(f'<text x="{x+40}" y="{y+40}" text-anchor="middle" font-size="{15 if len(label)>5 else 20}" font-weight="700" fill="{"#475569" if v is None or table_mode else "#ffffff"}" style="paint-order:stroke;stroke:{"transparent" if v is None or table_mode else "#334155"};stroke-width:.8px">{escape(label)}</text>')
    out.append('<text x="350" y="470" text-anchor="middle" font-size="15" fill="#334155">실측값 mils · 회색/—는 미측정'+('' if table_mode else ' · 음영은 입체 표현')+'</text>')
    return ''.join(out)+'</svg>'


def measurement_stats(grid):
    values=[v for row in grid for v in row if v is not None]
    return {'최소값':min(values) if values else None,'최댓값':max(values) if values else None,'Max−Min':store.flatness_span(grid),'측정 수':len(values)}


def select_measurement(kind,equipment,identifier,edit=False,part_type=None):
    suffix=f'_{part_type}' if part_type else ''
    st.session_state[f'measure_selected_{kind}_{equipment}{suffix}']=identifier
    if part_type:st.session_state[f'flatness_type_{equipment}']=part_type
    st.session_state['measurement_edit_requested']=identifier if edit else None


def measurement_history(rows,kind,equipment):
    spacing=kind=='spacingMeasurements'
    # 현재 설비·챔버로 필터된 기록만 전달받아 표를 만듭니다. 새로운 선택 창은 만들지 않습니다.
    groups={}
    for r in rows:groups.setdefault((r['chamberId'],'Spacing' if spacing else r['type']),[]).append(r)
    for (chamber,part),items in sorted(groups.items()):
        with st.container(border=True):
            st.subheader(f'{chamber} {part} 측정 이력 · {len(items)}건')
            st.caption(f'설비: {equipment} · 같은 챔버·같은 측정 종류의 이력만 표시합니다.')
            headers=['No','측정일','구분','부품 S/N','측정기','Max−Min','비고','보기','수정']
            widths=[.4,1.1,.8,1.4,1.1,.9,1.7,.55,.55]
            for c,label in zip(st.columns(widths),headers):c.caption(label)
            for n,r in enumerate(sorted(items,key=lambda item:(item['measuredAt'],item.get('originalOrder') or 0,item['id'])),1):
                grid=[r['values']] if spacing else r['grid']
                cells=st.columns(widths)
                values=[n,r['measuredAt'],r.get('timing',''),r.get('partSnAtMeasurement',r.get('partSn','')),r.get('gaugeTool',''),measurement_number(store.flatness_span(grid))+' mils',r.get('notes','')]
                for cell,value in zip(cells,values):cell.caption(str(value or '—'))
                cells[7].button('보기',key=f'measure_view_{kind}_{r["id"]}',on_click=select_measurement,args=(kind,equipment,r['id'],False,r.get('type')))
                cells[8].button('수정',key=f'measure_edit_{kind}_{r["id"]}',on_click=select_measurement,args=(kind,equipment,r['id'],True,r.get('type')))


def select_flatness_type(equipment,part):
    st.session_state[f'flatness_type_{equipment}']=part
    st.session_state.pop('measurement_edit_requested',None)


def flatness_type_cards(equipment,rows):
    key=f'flatness_type_{equipment}'
    selected=st.session_state.get(key,'Diffuser')
    if selected not in ['Diffuser','Susceptor']:selected='Diffuser'
    for col,part in zip(st.columns(2),['Diffuser','Susceptor']):
        with col:
            with st.container(border=True):
                count=sum(r.get('type')==part for r in rows)
                st.subheader(f'{part} Flatness')
                st.caption(f'현재 설비·챔버 범위의 측정 {count}건')
                st.button(f'{part} 데이터 보기',key=f'flatness_card_{equipment}_{part}',type='primary' if selected==part else 'secondary',on_click=select_flatness_type,args=(equipment,part),use_container_width=True)
    return selected


def comparison_view(rows,current,kind,equipment):
    spacing=kind=='spacingMeasurements'
    candidates=sorted([row for row in rows if row['chamberId']==current['chamberId'] and (spacing or row['type']==current['type'])],key=lambda row:(row['measuredAt'],row.get('originalOrder') or 0,row['id']))
    if len(candidates)<2:
        st.info('같은 설비·챔버·측정 종류에 기록이 2건 이상 있어야 전후 비교할 수 있습니다.');return
    indexed={row['id']:row for row in candidates};options=list(indexed)
    index=options.index(current['id']);after_index=max(1,index)
    context=f"{kind}_{equipment}_{current['chamberId']}_{current.get('type','Spacing')}_{current['id']}"
    def label(identifier):
        item=indexed[identifier]
        return f"{item['measuredAt']} · {item.get('timing','')} · {item.get('partSnAtMeasurement',item.get('partSn',''))} · {identifier}"
    left,right=st.columns(2)
    before_id=left.selectbox('비교 기준 A · 보통 작업 전',options,index=after_index-1,format_func=label,key='compare_A_'+context)
    after_id=right.selectbox('비교 대상 B · 보통 작업 후',options,index=after_index,format_func=label,key='compare_B_'+context)
    if before_id==after_id:st.info('전후 비교에는 서로 다른 두 기록을 선택해주세요.');return
    before=indexed[before_id];after=indexed[after_id]
    result=store.compare_measurements(before,after,kind)
    if before['measuredAt']>after['measuredAt']:st.warning('A 날짜가 B보다 늦습니다. 변화량은 선택한 순서대로 B−A로 계산합니다.')
    st.caption('전후 비교는 저장된 측정 기록 기준이며, 저장하지 않은 편집값은 포함하지 않습니다.')
    st.caption(f"비교 범위: {equipment} / {current['chamberId']} / {'Spacing' if spacing else current['type']} · 변화량 = B − A · 비교 가능 위치 {result['comparable']}/{9 if spacing else 35}")
    domain=(result['minimum'],result['maximum'])
    for col,item,caption in [(left,before,'A · 비교 기준'),(right,after,'B · 비교 대상')]:
        with col:
            st.subheader(caption)
            st.caption(f"{item['measuredAt']} · {item.get('timing','')} · {item.get('partSnAtMeasurement',item.get('partSn',''))}")
            grid=[item['values']] if spacing else item['grid']
            if spacing:
                st.markdown(spacing_diagram_svg(item['values']),unsafe_allow_html=True)
                value_table=pd.DataFrame([[item['values'][p-1] for p in row] for row in store.POINT_LAYOUT],columns=['좌','중앙','우'],index=['G/V','중앙','Window'])
            else:
                st.markdown(flatness_diagram_svg(item['grid'],color_range=domain),unsafe_allow_html=True)
                value_table=pd.DataFrame(item['grid'],columns=[f'C{i}' for i in range(1,8)],index=[f'R{i}' for i in range(1,6)])
            st.dataframe(value_table,use_container_width=True)
            stats=measurement_stats(grid)
            st.caption(f"최소 {measurement_number(stats['최소값'])} / 최대 {measurement_number(stats['최댓값'])} / Max−Min {measurement_number(stats['Max−Min'])} mils · 측정 {stats['측정 수']}개")
    if not spacing:
        legend='#64748b' if domain[0]==domain[1] else 'linear-gradient(90deg,#2563eb,#ef4444)'
        st.markdown(f'<div class="measure-legend" style="background:{legend}"></div>',unsafe_allow_html=True)
        st.caption(f'공통 색상 범위: 낮은 값 파랑 {measurement_number(domain[0])} → 높은 값 빨강 {measurement_number(domain[1])} mils · A/B에 같은 색상 기준 적용')
        if domain[0] is not None:
            with st.expander('전후 3D 비교 · 공통 색상 및 Z축 범위'):
                low,high=domain
                if low==high:low-=1;high+=1
                for col,item in zip(st.columns(2),[before,after]):
                    with col:
                        colors=[[0,'#64748b'],[1,'#64748b']] if domain[0]==domain[1] else [[0,MEASURE_BLUE],[1,MEASURE_RED]]
                        fig=go.Figure(go.Surface(z=item['grid'],x=list(range(1,8)),y=list(range(1,6)),colorscale=colors,cmin=low,cmax=high,connectgaps=False,lighting={'ambient':.65,'diffuse':.75,'specular':.18,'roughness':.85},colorbar={'title':'mils'}))
                        fig.update_layout(height=380,title=item['measuredAt'],scene={'xaxis_title':'Col','yaxis_title':'Row','zaxis':{'title':'mils','range':[low,high]}})
                        st.plotly_chart(fig,use_container_width=True)
    st.subheader('위치별 변화량 · B − A')
    difference=result['difference']
    if spacing:
        delta_values=difference[0]
        delta_grid=[[delta_values[p-1] for p in row] for row in store.POINT_LAYOUT]
        st.dataframe(pd.DataFrame(delta_grid,columns=['좌','중앙','우'],index=['G/V','중앙','Window']),use_container_width=True)
        st.subheader('Spacing 차이 항목 전후 비교')
        st.dataframe(pd.DataFrame(result['spacing_deltas']),hide_index=True,use_container_width=True)
        exports=[{'위치':f'P{i+1}','A 기록':before_id,'B 기록':after_id,'A(mils)':before['values'][i],'B(mils)':after['values'][i],'B−A(mils)':delta_values[i]} for i in range(9)]
    else:
        st.dataframe(pd.DataFrame(difference,columns=[f'C{i}' for i in range(1,8)],index=[f'R{i}' for i in range(1,6)]),use_container_width=True)
        exports=[{'위치':f'R{ri+1}C{ci+1}','A 기록':before_id,'B 기록':after_id,'A(mils)':before['grid'][ri][ci],'B(mils)':after['grid'][ri][ci],'B−A(mils)':difference[ri][ci]} for ri in range(5) for ci in range(7)]
    st.caption('양쪽 모두 실측값이 있는 위치만 차이를 계산합니다. 0은 유효값이며, 한쪽이 미측정이면 변화량도 빈값입니다. 비교는 저장된 기록을 읽기만 합니다.')
    export_buttons(pd.DataFrame(exports),'hw_compare_'+kind)


def measurement_export(rows,kind,spacing):
    export=pd.DataFrame([{'ID':row['id'],'설비':row['equipmentId'],'챔버':row['chamberId'],'측정일':row['measuredAt'],'종류':'Spacing' if spacing else row['type'],'측정값':json.dumps(row['values'] if spacing else row['grid']),'비고':row.get('notes','')} for row in rows])
    export_buttons(export,'hw_measure_'+kind)
