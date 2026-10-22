"""원본을 읽기만 하는 Excel 정규화 및 내보내기."""
from datetime import date, datetime
from io import BytesIO
import re
import pandas as pd
from openpyxl import load_workbook

MAIN = ['HW(설계)', '공정', 'SW', '전장', '기타']
DONE = {'O', '○', '완료'}
COLUMNS = ['ID','범위','분야','설비','번호','상태','원문 상태','Unit','내용','발생일','담당자','원인','진행 사항','목표일','조치내용','완료일','조치자','45호기 적용 여부','시트','원본 행','원문 전체']

def clean(value):
    if value is None: return ''
    if isinstance(value, (datetime,date)): return value.strftime('%Y-%m-%d')
    return str(value).strip()

def classify(*values):
    values = [clean(v).upper() for v in values]
    if all(v in DONE for v in values): return '완료'
    if any(not v or v in {'확인필요','확인 필요'} for v in values): return '확인 필요'
    return '미완료'

def equipment(value):
    text = clean(value).upper()
    found = [x for x in ['CVD15I','CVD45I'] if x in text]
    return ' / '.join(found) or '미기재'

def load_pending(content):
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    records = []
    supported = MAIN + ['고객사요청','software불합리','설비알람']
    for ws in wb:
        if ws.title not in supported: continue
        for n, row in enumerate(ws.iter_rows(values_only=True), 1):
            v = lambda col: clean(row[col-1]) if col <= len(row) else ''
            number_col = 2 if ws.title in MAIN+['고객사요청'] else 1
            if not re.fullmatch(r'\d+(?:\.0)?', v(number_col)): continue
            r = dict.fromkeys(COLUMNS, '')
            r.update({'ID':f'{ws.title}:{n}', '범위':'주요 분야' if ws.title in MAIN else '보조 시트', '분야':ws.title, '번호':v(number_col),'시트':ws.title,'원본 행':n})
            if ws.title in MAIN:
                sw = ws.title == 'SW'
                target, action, status, completed, owner, apply = (16,17,18,19,20,21) if sw else (17,18,19,20,21,22)
                mapping = {'설비':3,'발생일':5,'담당자':6,'Unit':7,'내용':8,'원인':11,'진행 사항':14,'목표일':target,'조치내용':action,'원문 상태':status,'완료일':completed,'조치자':owner,'45호기 적용 여부':apply}
                r.update({key:v(col) for key,col in mapping.items()})
                r['상태'] = classify(v(status))
            elif ws.title == '고객사요청':
                r.update({'Unit':v(4),'내용':v(5),'원인':v(6) if n<39 else '', '진행 사항':v(11),'45호기 적용 여부':v(13)})
                if n<39:
                    r.update({'설비':'미기재','원문 상태':f'H: {v(8)} / I: {v(9)}','상태':classify(v(8),v(9))})
                else:
                    r.update({'설비':v(7),'발생일':v(6),'목표일':v(8),'완료일':v(9),'원문 상태':v(10),'상태':classify(v(10))})
            elif ws.title == 'software불합리':
                r.update({'설비':v(2),'원문 상태':v(3),'상태':classify(v(3)),'목표일':v(4),'완료일':v(5),'내용':v(6),'진행 사항':v(7)})
            else:
                r.update({'설비':v(3),'발생일':v(2),'Unit':v(4),'내용':v(5)+'\n'+v(6),'원인':v(7),'조치내용':v(8),'진행 사항':v(9),'상태':'확인 필요'})
            if not r['내용']: continue
            r['설비'] = equipment(r['설비'])
            r['원문 전체'] = '\n'.join(f'{i}: {clean(x)}' for i,x in enumerate(row,1) if clean(x))
            records.append(r)
    wb.close()
    if not records: raise ValueError('지원하는 Pending 표를 찾지 못했습니다. 원본과 동일한 시트·열 구조인지 확인해주세요.')
    return pd.DataFrame(records)[COLUMNS]

def filter_pending(df, equipments, fields, states, query=''):
    mask = df['설비'].isin(equipments) & df['분야'].isin(fields) & df['상태'].isin(states)
    result = df.loc[mask]
    # 공백으로 나눈 모든 단어가 원문을 포함한 항목 전체에 있어야 합니다.
    for term in query.split():
        result = result[result.astype(str).agg(' '.join,axis=1).str.contains(term,case=False,regex=False)]
    return result

def summary(df, group):
    if df.empty: return pd.DataFrame(columns=[group,'전체','완료','미완료','확인 필요','완료율(%)'])
    t = pd.crosstab(df[group],df['상태']).reindex(columns=['완료','미완료','확인 필요'],fill_value=0)
    t['전체'] = t.sum(axis=1)
    t['완료율(%)'] = (t['완료']/t['전체']*100).round(1)
    return t.reset_index()

def safe_export(df):
    result = df.copy()
    for col in result.select_dtypes(include='object'):
        result[col] = result[col].map(lambda x: "'"+x if isinstance(x,str) and x.lstrip().startswith(('=','+','-','@')) else x)
    return result

def excel_bytes(df):
    buf = BytesIO()
    with pd.ExcelWriter(buf,engine='openpyxl') as writer:
        safe_export(df).to_excel(writer,index=False,sheet_name='조회 결과')
        summary(df,'분야').to_excel(writer,index=False,sheet_name='분야별 집계')
        summary(df,'설비').to_excel(writer,index=False,sheet_name='설비별 집계')
    return buf.getvalue()


def load_revision(content):
    """Rev 관리의 저장된 요약 값과 원문 이력. 수식 재계산은 하지 않습니다."""
    summary_columns=['구분','총 건수','조치 건수','조치율']
    history_columns=['Rev No.','구분','주요 변경내용','변경일자','작성자','변경점 Key Word / 비고','원본 행']
    wb=load_workbook(BytesIO(content),data_only=True,read_only=True)
    try:
        sheet=next((s for s in wb if re.sub(r'\s','',s.title).casefold()=='rev관리'),None)
        if sheet is None:
            return {'sheet':'','title':'','line':'','process':'','summary':pd.DataFrame(columns=summary_columns),'history':pd.DataFrame(columns=history_columns)}
        rows=list(sheet.iter_rows(values_only=True))
        def val(row,col):
            return rows[row-1][col-1] if row<=len(rows) and col<=len(rows[row-1]) else None
        totals=[]
        for col in [2,4,6,8,10,12,14]:
            label=clean(val(5,col))
            if not label: continue
            rate=val(8,col+1)
            rate=f'{rate*100:.1f}%' if isinstance(rate,(int,float)) else clean(rate) or '미기재 / 캐시 없음'
            totals.append({'구분':label,'총 건수':clean(val(7,col)) or '미기재 / 캐시 없음','조치 건수':clean(val(7,col+1)) or '미기재 / 캐시 없음','조치율':rate})
        history=[]
        for n in range(11,len(rows)+1):
            number=clean(val(n,2))
            if not re.fullmatch(r'\d+(?:\.0)?',number) or not clean(val(n,4)): continue
            history.append(dict(zip(history_columns,[number,clean(val(n,3)),clean(val(n,4)),clean(val(n,10)),clean(val(n,11)),clean(val(n,12)),n])))
        return {'sheet':sheet.title,'title':clean(val(2,4)),'line':clean(val(2,13)),'process':clean(val(3,13)), 'summary':pd.DataFrame(totals,columns=summary_columns),'history':pd.DataFrame(history,columns=history_columns)}
    finally:
        wb.close()


def revision_display(revisions):
    """조회용 정리. 원본/DB를 수정하지 않고 전후 내역을 반환합니다."""
    result=revisions.copy(deep=True)
    changes=[]
    for idx,row in revisions.iterrows():
        for field in ['구분','주요 변경내용','변경일자','작성자','변경점 Key Word / 비고']:
            before=clean(row[field]);after=before;reason=[]
            if field=='변경일자':
                parts=[p.strip() for p in re.split(r'(?:\r?\n|\\n|/n|\s+/\s+)+',before) if p.strip()]
                unique=[];seen=set()
                for part in parts:
                    token=part
                    match=re.fullmatch(r'(\d{4})[-.](\d{1,2})[-.](\d{1,2})',part)
                    if match: token=tuple(int(x) for x in match.groups())
                    if token not in seen:
                        unique.append(part);seen.add(token)
                if len(unique)<len(parts):
                    after=' / '.join(unique);reason.append('동일 변경일자 중복 제거')
                elif len(parts)>1: after=' / '.join(parts)
            if field=='작성자' and ('\n' in after or '\r' in after):
                after=' / '.join(part.strip() for part in after.splitlines() if part.strip())
                reason.append('작성자 줄바꿈을 / 구분으로 표시')
            if '/n' in after or '\\n' in after:
                after=after.replace('/n','/').replace('\\n','\\')
                reason.append('문자열 줄바꿈 표기의 n만 제거, 슬래시 보존')
            if after!=before:
                result.at[idx,field]=after
                changes.append({'REV':clean(row['Rev No.']),'원본 행':row.get('원본 행',''),'필드':field,'정리 전':before,'정리 후':after,'사유':' / '.join(reason) or '여러 날짜 구분 표시'})
    return result,pd.DataFrame(changes,columns=['REV','원본 행','필드','정리 전','정리 후','사유'])


def revision_table_html(df):
    """pandas.to_html의 줄바꿈 이스케이프를 피하고 안전하게 셀을 렌더링합니다."""
    from html import escape
    columns=[c for c in df.columns if c not in ['원본 행','ID','버전']]
    def cell(value):
        return escape(str(value)).replace('\r\n','\n').replace('\r','\n').replace('\n','<br>')
    head='<thead><tr>'+''.join('<th>'+escape(c)+'</th>' for c in columns)+'</tr></thead>'
    rows=''.join('<tr>'+''.join('<td>'+cell(row[c])+'</td>' for c in columns)+'</tr>' for _,row in df.iterrows())
    return '<table class="rev-history-table">'+head+'<tbody>'+rows+'</tbody></table>'
