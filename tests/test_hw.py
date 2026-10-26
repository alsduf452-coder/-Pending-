import base64
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import shutil
from zipfile import ZipFile
import pytest
from streamlit.testing.v1 import AppTest
import hw_store as hw
from hw_ui import full_backup


def fixture():
    return {'mode':'sample','equipments':[{'id':'A4CVD15I','name':'설비15','chambers':[{'chamberId':'Chamber A','initialParts':{'Diffuser':{'sn':'OLD','installedDate':'2025-01-01'}},'installedParts':{}}]}],
            'workLogs':[{'id':'WL1','equipmentId':'A4CVD15I','chamberId':'Chamber A','workDate':'2025-02-01','workType':'부품 교체','details':'부품 교체','beforeSn':'OLD','afterSn':'NEW','targetPart':'Diffuser','attachments':None}],
            'spacingMeasurements':[{'id':'SP1','equipmentId':'A4CVD15I','chamberId':'Chamber A','measuredAt':'2025-02-01','values':[0,None,2,3,4,5,6,7,8],'unit':'mils','relatedWorkLogId':'WL1'}],
            'flatnessMeasurements':[{'id':'FL1','equipmentId':'A4CVD15I','chamberId':'Chamber A','measuredAt':'2025-02-01','type':'Diffuser','grid':[[0,None,2,3,4,5,6]]*5,'unit':'mils'}],
            'spacingStandards':[{'pointIndex':i,'label':f'P{i}'} for i in range(1,10)]}


def encoded(data):return json.dumps(data,ensure_ascii=False).encode()


def test_import_validation_dedup_and_original_values(tmp_path):
    raw=fixture();raw['workLogs'][0]['details']=''
    plan=hw.parse_import(encoded(raw))
    assert not plan['errors'] and plan['warnings']
    db=tmp_path/'hw.sqlite3'
    result=hw.import_backup(db,encoded(raw),'등록자')
    assert result['added']==13
    assert hw.import_backup(db,encoded(raw),'등록자')['added']==0
    assert hw.records(db,'workLogs')[0]['details']==''
    assert hw.records(db,'spacingMeasurements')[0]['values'][:2]==[0,None]
    assert hw.records(db,'flatnessMeasurements')[0]['grid'][0][:2]==[0,None]
    invalid=deepcopy(raw);invalid['spacingMeasurements'][0]['values']=[None]
    assert hw.parse_import(encoded(invalid))['errors']
    with pytest.raises(ValueError):hw.import_backup(tmp_path/'bad.sqlite3',encoded(invalid),'人')
    assert hw.records(tmp_path/'bad.sqlite3','equipments')==[]
    invalid['workLogs'].append(deepcopy(invalid['workLogs'][0]))
    assert any('중복' in r['사유'] for r in hw.parse_import(encoded(invalid))['errors'])
    invalid=deepcopy(raw);invalid['flatnessMeasurements'][0]['chamberId']='없는챔버'
    assert hw.parse_import(encoded(invalid))['errors']
    invalid=deepcopy(raw);invalid['flatnessMeasurements'][0].pop('grid');invalid['flatnessMeasurements'][0]['gridRows']=[{'cells':x} for x in raw['flatnessMeasurements'][0]['grid']]
    assert not hw.parse_import(encoded(invalid))['errors']


def test_edit_concurrency_attachments_backup_and_calculations(tmp_path):
    db=tmp_path/'hw.sqlite3';hw.import_backup(db,encoded(fixture()),'가져오기')
    current=hw.records(db,'workLogs')[0]
    updated={**current,'details':'수정 내용','pendingLink':{'source':'excel-hash','id':'SW:6','title':'관련 이슈'}}
    raw=b'file bytes'
    hw.save_record(db,'workLogs',updated,'수정자',current['_version'],[{'name':'check.pdf','mime':'application/pdf','data':raw}])
    with pytest.raises(ValueError,match='다른 화면'):hw.save_record(db,'workLogs',updated,'다른 사용자',current['_version'])
    saved=hw.records(db,'workLogs')[0]
    assert hw.get_blob(db,saved['attachments'][0]['blobId'])['data']==raw
    restored=tmp_path/'restored.sqlite3'
    hw.import_backup(restored,hw.backup_bytes(db),'복원자')
    assert hw.records(restored,'workLogs')[0]['pendingLink']['source']=='excel-hash'
    assert hw.get_blob(restored,saved['attachments'][0]['blobId'])['data']==raw
    assert any(r['actor']=='수정자' for r in hw.audit_records(restored))
    assert hw.spacing_deltas([0,None,2,3,4,5,6,7,8])['5−1']==4
    assert hw.spacing_deltas([0,None,2,3,4,5,6,7,8])['5−2'] is None
    assert hw.flatness_span([[0,None],[2,-1]])==3
    eq=hw.records(db,'equipments')[0]
    assert hw.installed_parts(eq,'Chamber A',[saved])[0]['sn']=='NEW'
    assert any(r['구분']=='장착 후 기록' for r in hw.trace_sn([eq],[saved],'NEW'))
    data_dir=tmp_path/'data';data_dir.mkdir();shutil.copyfile(db,data_dir/'hw_history.sqlite3')
    archive=full_backup(data_dir)
    with ZipFile(BytesIO(archive)) as z:
        assert 'data/hw_history.sqlite3' in z.namelist()
        assert 'RESTORE.txt' in z.namelist()


def setup_ui(tmp_path):
    root=Path(__file__).resolve().parents[1]
    for name in ['app.py','pending_data.py','pending_store.py','hw_ui.py','hw_store.py']:
        shutil.copyfile(root/name,tmp_path/name)
    (tmp_path/'data').mkdir()
    return tmp_path/'data/hw_history.sqlite3'


def test_hw_without_pending_excel_and_all_screens(tmp_path):
    db=setup_ui(tmp_path)
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='nav_main_page_'+'설비 H/W 이력관리').click().run()
    assert not app.exception
    assert any('아직 H/W 데이터' in r.value for r in app.info)
    hw.import_backup(db,encoded(fixture()),'준비')
    app.run()
    assert not app.exception
    for section in ['작업 이력','부품 추적','Spacing','Flatness','데이터 관리','설비 현황']:
        app.button(key='nav_hw_section_'+section).click().run()
        assert not app.exception,section


def test_hw_work_registration_ui_and_audit(tmp_path):
    db=setup_ui(tmp_path);hw.import_backup(db,encoded(fixture()),'준비')
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='nav_main_page_'+'설비 H/W 이력관리').click().run()
    app.button(key='nav_hw_section_'+'작업 이력').click().run()
    # New form precedes the edit form. Fill required inputs in the new form.
    next(x for x in app.text_area if x.label=='작업 내용 *').set_value('화면 신규 점검')
    next(x for x in app.text_input if x.label=='등록·변경자 *').set_value('작업자')
    next(x for x in app.button if x.label=='작업 이력 저장').click().run()
    assert not app.exception
    assert len(hw.records(db,'workLogs'))==2
    assert any(r['actor']=='작업자' for r in hw.audit_records(db))


def test_pending_work_link_roundtrip(tmp_path):
    from test_pending import workbook
    db=setup_ui(tmp_path)
    source=tmp_path/'data/default.xlsx';source.write_bytes(workbook())
    data=fixture()
    data['equipments'][0]['id']='A4CVD45I'
    for kind in ['workLogs','spacingMeasurements','flatnessMeasurements']:
        for r in data[kind]:r['equipmentId']='A4CVD45I'
    hw.import_backup(db,encoded(data),'준비')
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='pending_to_hw').click().run()
    assert not app.exception
    assert app.session_state['main_page']=='설비 H/W 이력관리'
    assert app.session_state['hw_section']=='작업 이력'
    next(x for x in app.text_area if x.label=='작업 내용 *').set_value('Pending 관련 교체')
    next(x for x in app.text_input if x.label=='등록·변경자 *').set_value('연결 작업자')
    next(x for x in app.button if x.label=='작업 이력 저장').click().run()
    saved=hw.records(db,'workLogs')[-1]
    assert saved['pendingLink']['id']=='SW:6'
    import hashlib
    assert saved['pendingLink']['source']==hashlib.sha256(source.read_bytes()).hexdigest()
    selector=next(x for x in app.selectbox if x.label=='작업 상세 선택')
    selector.set_value(saved['id']).run()
    app.button(key='hw_back_pending').click().run()
    assert not app.exception
    assert app.session_state['main_page']=='PENDING 현황'
    assert app.text_input(key='keyword').value=='SW:6'
    assert any('관련 H/W 작업 이력' in x.label for x in app.expander)


def test_age_and_monthly_filters():
    from datetime import date
    today=date(2026,10,22)
    assert hw.installation_age('2025-10-22','Diffuser',today)['days']==365
    assert '#fee2e2' in hw.installation_age('2025-10-22','Diffuser',today)['style']
    assert hw.installation_age('2025-10-23','Diffuser',today)['days']==364
    assert hw.installation_age('2020-01-01','RFG',today)['days'] is None
    assert hw.installation_age('no date','Diffuser',today)['days'] is None
    assert '미래' in hw.installation_age('2027-01-01','Diffuser',today)['label']
    logs=fixture()['workLogs']
    logs.append({**logs[0],'id':'setup','workType':'초기 셋업','workDate':'2025-01-01'})
    logs.append({**logs[0],'id':'other','workDate':'2026-01-01','targetPart':'Matcher'})
    counts,rows=hw.monthly_work(logs,'A4CVD15I',2025)
    assert len(counts)==12 and len(rows)==1
    assert sum(sum(v for k,v in r.items() if k!='월') for r in counts)==len(rows)
    assert len(hw.monthly_work(logs,'A4CVD15I',2025,mode='전체 작업')[1])==2
    assert not hw.monthly_work(logs,'A4CVD15I',2026,part='Diffuser')[1]
    assert len(hw.monthly_work(logs,'A4CVD15I',2026,part='Matcher')[1])==1
    assert not hw.monthly_work(logs,'A4CVD15I',2025,chamber='Chamber B')[1]
    assert not hw.is_replacement(logs[1])


def test_archive_restore_and_backup_retains_history(tmp_path):
    path=tmp_path/'hw.sqlite3';hw.import_backup(path,encoded(fixture()),'등록자')
    row=hw.records(path,'workLogs')[0]
    with pytest.raises(ValueError):hw.archive_work(path,row['id'],'담당자','',row['_version'])
    hw.archive_work(path,row['id'],'삭제자','중복 정리',row['_version'])
    assert hw.records(path,'workLogs')==[]
    archived=hw.records(path,'workLogs',include_deleted=True)[0]
    assert archived['deletionReason']=='중복 정리'
    assert len(hw.records(path,'spacingMeasurements'))==1
    restored=tmp_path/'restored.sqlite3';hw.import_backup(restored,hw.backup_bytes(path),'백업복원')
    assert hw.records(restored,'workLogs')==[]
    with pytest.raises(ValueError):hw.archive_work(path,row['id'],'복원자','복원',row['_version'],restore=True)
    hw.archive_work(path,row['id'],'복원자','이력 복구',archived['_version'],restore=True)
    assert len(hw.records(path,'workLogs'))==1
    assert hw.records(path,'workLogs')[0]['restorationReason']=='이력 복구'
    assert {'IMPORT','ARCHIVE','RESTORE'}.issubset({r['action'] for r in hw.audit_records(path)})


def test_monthly_controls_and_card_actions(tmp_path):
    db=setup_ui(tmp_path);data=fixture()
    data['equipments'][0]['chambers'].append({'chamberId':'Chamber B','initialParts':{'Diffuser':{'sn':'B-SN','installedDate':'2025-02-01'}},'installedParts':{}})
    hw.import_backup(db,encoded(data),'등록자')
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='nav_main_page_'+'설비 H/W 이력관리').click().run()
    assert not app.exception
    assert sum('<table class="hw-parts-table">' in x.value for x in app.markdown)==1
    assert any('<th scope="col">Chamber A</th>' in x.value and 'Chamber A' in x.value and 'Chamber B' in x.value and 'B-SN' in x.value for x in app.markdown)
    assert app.selectbox(key='monthly_part_A4CVD15I').options[1:9]==hw.PARTS[:8]
    app.selectbox(key='monthly_year_A4CVD15I').set_value(2025).run()
    app.selectbox(key='monthly_chamber_A4CVD15I').set_value('전체 Chamber').run()
    assert not app.exception
    app.button(key='nav_hw_section_'+'작업 이력').click().run()
    app.button(key='card_edit_WL1').click().run()
    assert not app.exception and app.session_state['hw_edit_id']=='WL1'
    app.button(key='card_trace_WL1').click().run()
    assert not app.exception and app.session_state['hw_section']=='부품 추적'
    assert app.text_input(key='hw_trace_sn').value=='NEW'
    app.button(key='nav_hw_section_'+'작업 이력').click().run()
    app.button(key='card_archive_WL1').click().run()
    next(x for x in app.text_input if x.label=='삭제 담당자 *').set_value('삭제자')
    next(x for x in app.text_input if x.label=='삭제 사유 *').set_value('중복')
    next(x for x in app.button if x.label=='삭제 보관 확인').click().run()
    assert not app.exception and not hw.records(db,'workLogs')
    assert not any('삭제 보관함' in x.label for x in app.expander)
    archived=hw.records(db,'workLogs',include_deleted=True)[0]
    assert archived['deletedAt']



def test_measurement_diagrams_and_color_direction():
    from hw_ui import flatness_color, flatness_diagram_svg, spacing_diagram_svg, spacing_table_html, flatness_table_html, MEASURE_BLUE, MEASURE_RED
    assert flatness_color(-40,-40,19)==MEASURE_BLUE
    assert flatness_color(19,-40,19)==MEASURE_RED
    assert flatness_color(None,-40,19)=='#e2e8f0'
    assert flatness_color(0,0,0)=='#64748b'
    grid=[[-40,19,0,None,1,2,3]]*5
    numeric=flatness_table_html(grid)
    assert '<table' in numeric and '<svg' not in numeric
    assert 'Row 5' in numeric and 'Col 7' in numeric and 'R5C7' in numeric
    assert '>0<' in numeric and '미측정' in numeric
    svg=flatness_diagram_svg(grid)
    assert 'R1C1' in svg and 'R5C7' in svg and 'flatshade' in svg
    assert MEASURE_BLUE in svg and MEASURE_RED in svg
    assert '>0<' in svg and '>—<' in svg
    spacing=[867.2,None,886.4,888.1,885.1,872.6,883.6,890.0,884.7]
    html=spacing_table_html(spacing);diagram=spacing_diagram_svg(spacing)
    assert '867.2' in html and '미측정' in html
    assert 'GATE VALVE' in diagram and 'WINDOW' in diagram
    assert 'P6 · G/V 중앙' in diagram and 'P9 · 중앙 좌' in diagram
    assert 'RdBu' not in svg+diagram


def test_measurement_split_ui_and_chamber_history(tmp_path):
    db=setup_ui(tmp_path);data=fixture()
    data['equipments'][0]['chambers'].append({'chamberId':'Chamber B','initialParts':{},'installedParts':{}})
    b={**data['spacingMeasurements'][0],'id':'SP_B','chamberId':'Chamber B','relatedWorkLogId':None}
    data['spacingMeasurements'].append(b)
    hw.import_backup(db,encoded(data),'등록자')
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='nav_main_page_'+'설비 H/W 이력관리').click().run()
    app.selectbox(key='hw_chamber_A4CVD15I').set_value('Chamber B').run()
    app.button(key='nav_hw_section_'+'Spacing').click().run()
    assert not app.exception
    headers=[x.value for x in app.subheader]
    assert any('실측 데이터' in x for x in headers) and any('측정 위치도' in x for x in headers)
    assert not any('50%' in x for x in headers)
    histories=[x for x in headers if '측정 이력 ·' in x]
    assert histories and all('Chamber B' in x for x in histories)
    assert not any(x.label=='측정 기록 선택' for x in app.selectbox)
    app.button(key='nav_hw_section_'+'Flatness').click().run()
    assert not app.exception
    app.selectbox(key='hw_chamber_A4CVD15I').set_value('Chamber A').run()
    assert not app.exception
    assert any('낮은 값 · 파랑' in x.value and '높은 값 · 빨강' in x.value for x in app.caption)
    assert any('Chamber A Diffuser 측정 이력' in x.value for x in app.subheader)
    key='preview_flatnessMeasurements_FL1_1'
    app.session_state[key]={'edited_rows':{0:{'Col 1':-10.0}},'added_rows':[],'deleted_rows':[]}
    app.run()
    assert not app.exception
    next(x for x in app.text_input if x.label=='수치 변경자 *').set_value('수치 편집자')
    next(x for x in app.button if x.label=='편집한 측정값 저장').click().run()
    assert not app.exception
    assert hw.records(db,'flatnessMeasurements')[0]['grid'][0][0]==-10
    assert hw.records(db,'flatnessMeasurements')[0]['grid'][0][1] is None


def test_measurement_comparison_values_and_scope():
    data=fixture();before=data['spacingMeasurements'][0]
    after={**before,'id':'SP2','values':[1,0,5,3,2,8,6,7,None]}
    result=hw.compare_measurements(before,after,'spacingMeasurements')
    assert result['difference'][0][:3]==[1,None,3]
    assert result['difference'][0][-1] is None
    assert result['spacing_deltas'][0]['변화 B−A(mils)']==-3
    assert result['comparable']==7
    with pytest.raises(ValueError):hw.compare_measurements(before,{**after,'chamberId':'Chamber B'},'spacingMeasurements')
    flat=data['flatnessMeasurements'][0]
    another={**flat,'id':'FL2','grid':[[10,None,-5,3,4,5,6]]*5}
    comparison=hw.compare_measurements(flat,another,'flatnessMeasurements')
    assert comparison['minimum']==-5 and comparison['maximum']==10
    assert comparison['difference'][0][:3]==[10,None,-7]
    with pytest.raises(ValueError):hw.compare_measurements(flat,{**another,'type':'Susceptor'},'flatnessMeasurements')
    assert flat['grid'][0][0]==0
    from hw_ui import flatness_diagram_svg
    # A common domain must yield the same colour for an equal number in different records.
    a=flatness_diagram_svg(flat['grid'],color_range=(-5,10))
    b=flatness_diagram_svg(another['grid'],color_range=(-5,10))
    assert 'flatshade' in a+b


def test_flatness_type_cards_and_comparison_ui(tmp_path):
    db=setup_ui(tmp_path);data=fixture()
    data['flatnessMeasurements'].append({**data['flatnessMeasurements'][0],'id':'FL2','measuredAt':'2025-03-01','grid':[[10,None,-5,3,4,5,6]]*5})
    data['flatnessMeasurements'].append({**data['flatnessMeasurements'][0],'id':'SUS1','type':'Susceptor','partSn':'SUS-SN','measuredAt':'2025-02-01'})
    data['flatnessMeasurements'].append({**data['flatnessMeasurements'][0],'id':'SUS2','type':'Susceptor','partSn':'SUS-SN','measuredAt':'2025-03-01'})
    data['spacingMeasurements'].append({**data['spacingMeasurements'][0],'id':'SP2','measuredAt':'2025-03-01'})
    hw.import_backup(db,encoded(data),'등록자')
    versions_before={kind:[(r['id'],r['_version']) for r in hw.records(db,kind)] for kind in hw.KINDS}
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='nav_main_page_'+'설비 H/W 이력관리').click().run()
    app.button(key='nav_hw_section_'+'Flatness').click().run()
    app.selectbox(key='hw_chamber_A4CVD15I').set_value('Chamber A').run()
    assert not app.exception
    assert any(x.label=='Diffuser 데이터 보기' for x in app.button)
    assert any(x.label=='Susceptor 데이터 보기' for x in app.button)
    app.button(key='flatness_card_A4CVD15I_Susceptor').click().run()
    assert not app.exception
    assert any('Susceptor Flatness 실측 데이터' in x.value for x in app.subheader)
    histories=[x.value for x in app.subheader if '측정 이력 ·' in x.value]
    assert histories and all('Susceptor' in x for x in histories)
    assert not any('50%' in x.value for x in app.subheader)
    assert any(x.value.startswith('측정 구분:') for x in app.caption)
    mode=next(x for x in app.radio if x.label=='측정 보기 방식')
    mode.set_value('작업 전후 비교').run()
    assert not app.exception
    options=next(x for x in app.selectbox if x.label=='비교 기준 A · 보통 작업 전').options
    assert all('SUS' in x for x in options) and not any('FL2' in x for x in options)
    assert any('공통 색상 범위' in x.value for x in app.caption)
    app.button(key='flatness_card_A4CVD15I_Diffuser').click().run()
    assert not app.exception
    next(x for x in app.radio if x.label=='측정 보기 방식').set_value('작업 전후 비교').run()
    assert not app.exception
    app.button(key='nav_hw_section_'+'Spacing').click().run()
    next(x for x in app.radio if x.label=='측정 보기 방식').set_value('작업 전후 비교').run()
    assert not app.exception
    assert any(x.value=='Spacing 차이 항목 전후 비교' for x in app.subheader)
    versions_after={kind:[(r['id'],r['_version']) for r in hw.records(db,kind)] for kind in hw.KINDS}
    assert versions_before==versions_after


def test_hw_excel_sheets_values_and_archived_work(tmp_path):
    from openpyxl import load_workbook
    from hw_ui import hw_excel_bytes
    db=tmp_path/'hw.sqlite3';data=fixture()
    data['workLogs'][0]['reason']='=1+1'
    data['flatnessMeasurements'].append({**data['flatnessMeasurements'][0],'id':'FS1','type':'Susceptor'})
    hw.import_backup(db,encoded(data),'검토자')
    book=load_workbook(BytesIO(hw_excel_bytes(db)))
    assert len(book.sheetnames)==11
    assert {'Spacing','Diffuser Flatness','Susceptor Flatness','작업 교체 이력','장착 부품','변경 이력'}.issubset(book.sheetnames)
    def value(sheet,column,row=2):
        ws=book[sheet];headers=[c.value for c in ws[1]]
        return ws.cell(row,headers.index(column)+1).value
    assert book['Spacing']['B13'].value==0
    assert book['Spacing']['D13'].value is None
    assert book['Diffuser Flatness']['B13'].value==0
    assert book['Diffuser Flatness']['C13'].value is None
    assert book['Susceptor Flatness']['B5'].value=='FS1'
    assert value('작업 교체 이력','사유')=="'=1+1"
    assert book['Spacing'].freeze_panes=='B4'
    assert '챔버 정보' not in book.sheetnames
    assert not {'관련 작업 ID','상태','계산 근거'} & {c.value for c in book['장착 부품'][1]}
    assert '원문 출처' not in {c.value for c in book['작업 교체 이력'][1]}
    original=hw.records(db,'workLogs')[0]
    hw.archive_work(db,'WL1','검토자','중복',original['_version'])
    book=load_workbook(BytesIO(hw_excel_bytes(db)))
    assert book['작업 교체 이력'].max_row==1
    assert book['삭제 보관 작업'].max_row==2
    assert value('장착 부품','S/N')=='OLD'
    empty=load_workbook(BytesIO(hw_excel_bytes(tmp_path/'empty.sqlite3')))
    assert len(empty.sheetnames)==11 and empty['Spacing']['A4'].value=='등록된 측정 데이터 없음'



def test_measurement_summary_scope_latest_and_months():
    from hw_ui import measurement_summary_data
    from copy import deepcopy
    first=fixture()['spacingMeasurements'][0]
    latest={**first,'id':'SP2','measuredAt':'2026-01-02','values':[None]*4+[0]+[None]*4}
    same_day={**first,'id':'SP3','measuredAt':'2026-01-02','originalOrder':2}
    rows=[first,latest,same_day,{**first,'id':'OLD','measuredAt':'2024-12-31'}, {**first,'id':'OTHER','equipmentId':'OTHER'}, {**first,'id':'DELETED','deletedAt':'2026-01-03'}]
    before=deepcopy(rows)
    summary,months=measurement_summary_data(rows,'A4CVD15I',['Chamber A','Chamber B'],'spacingMeasurements',year=2026)
    assert summary[0]['최신 측정일']=='2026-01-02'
    assert summary[0]['Center 값']==4
    assert summary[0]['측정 개수']=='8/9'
    assert summary[1]['최신 측정일']=='측정 기록 없음'
    assert summary[1]['Center 값'] is None
    assert len(months)==12 and months[0]['Chamber A']==2
    assert sum(row['Chamber A'] for row in months)==2
    zero,_=measurement_summary_data([latest],'A4CVD15I',['Chamber A'],'spacingMeasurements')
    assert zero[0]['Center 값']==0 and zero[0]['Max−Min']==0
    assert zero[0]['측정 개수']=='1/9'
    flat=fixture()['flatnessMeasurements'][0]
    flat_rows=[flat,{**flat,'id':'SUS1','type':'Susceptor','measuredAt':'2026-04-01'}]
    flat_summary,flat_months=measurement_summary_data(flat_rows,'A4CVD15I',['Chamber A'],'flatnessMeasurements','Diffuser',2026)
    assert flat_summary[0]['최신 측정일']=='2025-02-01'
    assert flat_summary[0]['최솟값']==0
    assert sum(r['Chamber A'] for r in flat_months)==0
    assert rows==before


def test_measurement_overview_ui_and_chamber_detail(tmp_path):
    db=setup_ui(tmp_path);data=fixture()
    data['equipments'][0]['chambers'].append({'chamberId':'Chamber B','initialParts':{},'installedParts':{}})
    data['spacingMeasurements'].extend([{**data['spacingMeasurements'][0],'id':'SP2','measuredAt':'2025-02-01'}, {**data['spacingMeasurements'][0],'id':'SP_OLD','measuredAt':'2024-01-01'}])
    data['flatnessMeasurements'].append({**data['flatnessMeasurements'][0],'id':'SUS1','type':'Susceptor','measuredAt':'2026-04-01'})
    hw.import_backup(db,encoded(data),'등록자')
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.button(key='nav_main_page_설비 H/W 이력관리').click().run()
    app.button(key='nav_hw_section_Spacing').click().run()
    assert not app.exception
    assert [x.value for x in app.subheader]==['챔버별 최신 측정 요약','월별 측정 횟수']
    assert not any('수치 직접 편집' in x.label for x in app.expander)
    assert not any(x.label in ['측정 보기 방식','측정 기록 검색'] for x in list(app.radio)+list(app.text_input))
    assert not any('새 측정 등록' in x.label for x in app.expander)
    year=next(x for x in app.selectbox if x.label=='측정 집계 연도')
    assert year.value==2025
    assert any('총 2회' in x.value for x in app.caption)
    year.set_value(2024).run()
    assert not app.exception
    summary=next(x.value for x in app.dataframe if '최신 측정일' in x.value.columns)
    assert summary.iloc[0]['최신 측정일']=='2025-02-01'
    assert summary.iloc[1]['최신 측정일']=='측정 기록 없음'
    assert any('총 1회' in x.value for x in app.caption)
    app.button(key='nav_hw_section_Flatness').click().run()
    assert not app.exception and any('Diffuser Flatness' in x.value for x in app.caption)
    app.button(key='flatness_card_A4CVD15I_Susceptor').click().run()
    assert not app.exception and any('Susceptor Flatness' in x.value for x in app.caption)
    summary=next(x.value for x in app.dataframe if '최신 측정일' in x.value.columns)
    assert summary.iloc[0]['최신 측정일']=='2026-04-01'
    app.selectbox(key='hw_chamber_A4CVD15I').set_value('Chamber A').run()
    assert not app.exception
    assert any('수치 직접 편집' in x.label for x in app.expander)
    assert any('Susceptor Flatness 실측 데이터' in x.value for x in app.subheader)
