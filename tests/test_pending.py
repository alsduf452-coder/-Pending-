from io import BytesIO
from pathlib import Path
from hashlib import sha256
import pandas as pd
from openpyxl import Workbook, load_workbook
from pending_data import *

def workbook():
    w=Workbook(); s=w.active; s.title='SW'
    s.cell(5,2,'No.')
    for col,value in {2:1,3:'A4CVD45I',8:'Robot Leak',16:'12월 24일',17:'조치 원문',18:'O'}.items(): s.cell(6,col,value)
    c=w.create_sheet('고객사요청')
    for row in [3,4,40]: c.cell(row,2,row); c.cell(row,5,'요청')
    c.cell(3,8,'완료'); c.cell(3,9,'완료')
    c.cell(4,8,'완료'); c.cell(4,9,'검토')
    c.cell(40,7,'A4CVD15I'); c.cell(40,10,'○')
    b=BytesIO();w.save(b);return b.getvalue()

def test_layout_and_status():
    d=load_pending(workbook())
    assert len(d)==4
    r=d.iloc[0];assert r['상태']=='완료' and r['조치내용']=='조치 원문' and r['목표일']=='12월 24일'
    assert d['상태'].tolist()==['완료','완료','미완료','완료']
    assert classify('')=='미완료'
    assert classify('X')=='미완료'

def test_filter_export():
    d=load_pending(workbook())
    r=filter_pending(d,['CVD45I'],['SW'],['완료'],'robot Leak')
    assert len(r)==1
    assert filter_pending(d,[],['SW'],['완료']).empty
    assert filter_pending(d,['CVD45I'],['SW'],['완료'],'[').empty
    x=load_workbook(BytesIO(excel_bytes(r)))
    assert len(x.sheetnames)==3 and x.worksheets[0].max_row==2
    assert safe_export(pd.DataFrame({'a':['=1+2']})).iloc[0,0]=="'=1+2"

def test_real_source_read_only():
    p=Path('data/source.xlsx')
    if not p.exists():
        import pytest;pytest.skip('원본 파일이 로컬에 없으므로 통합 테스트 생략')
    before=sha256(p.read_bytes()).hexdigest()
    d=load_pending(p.read_bytes())
    assert set(MAIN).issubset(d['분야'].unique())
    assert len(d[d['범위']=='주요 분야'])==174
    assert d['ID'].is_unique
    assert before==sha256(p.read_bytes()).hexdigest()

def test_revision():
    w=Workbook();s=w.active;s.title='REV 관리'
    s.cell(2,4,'통합 이력');s.cell(2,13,'라인A');s.cell(3,13,'공정A')
    s.cell(5,2,'Total');s.cell(7,2,10);s.cell(7,3,8);s.cell(8,3,0.8)
    s.cell(12,2,'000');s.cell(12,4,'첫째 줄\n둘째 줄');s.cell(12,10,'2026.10.16');s.cell(12,11,'작성자')
    b=BytesIO();w.save(b)
    rev=load_revision(b.getvalue())
    assert rev['line']=='라인A' and rev['summary'].iloc[0]['조치율']=='80.0%'
    assert rev['history'].iloc[0]['Rev No.']=='000'
    assert rev['history'].iloc[0]['주요 변경내용']=='첫째 줄\n둘째 줄'
    assert load_revision(workbook())['history'].empty

def test_real_revision():
    p=Path('data/source.xlsx')
    if not p.exists():
        import pytest;pytest.skip('원본 없는 환경')
    rev=load_revision(p.read_bytes())
    assert len(rev['history'])==18
    assert rev['summary'].iloc[0]['총 건수']=='174'

def test_binary_summary_and_drilldown():
    data=pd.DataFrame({'상태':['완료','확인 필요','미완료',''], '분야':['SW','SW','공정','공정'],'설비':['CVD15I']*4})
    normalized=binary_status(data)
    result=summary(data,'분야')
    assert '확인 필요' not in result.columns
    assert result['전체'].sum()==4 and result['미완료'].sum()==3
    assert len(drilldown_pending(normalized,{'group':'분야','value':'SW','state':'미완료'}))==1
    assert len(drilldown_pending(normalized,{'state':'완료'}))==1
    assert len(drilldown_pending(normalized,{'group':'분야','value':'없는분야','state':'미완료'}))==0



def test_pending_deadlines_boundaries_and_invalid_values():
    from datetime import date
    from pending_data import pending_deadlines
    df=pd.DataFrame({'상태':['미완료']*10+['완료'],'목표일':['2026-10-25','2026-10-26','2026-11-02','2026-11-03','26.10.25','2026/11/03','','2026-02-30','미정','2026-10-26 ~ 2026-10-30','2026-10-25']})
    before=df.copy(deep=True)
    result=pending_deadlines(df,date(2026,10,26))
    assert result['목표일 구분'].tolist()==['기한초과','임박','임박','여유','기한초과','여유']+['목표일 확인 대상']*4
    pd.testing.assert_frame_equal(df,before)
