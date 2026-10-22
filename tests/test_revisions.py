from io import BytesIO
from pathlib import Path
import shutil
import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from pending_data import load_revision
from pending_store import (read_revisions, save_revision, read_revision_history,
                           suggest_revision, read_items, source_key)
from test_pending import workbook


def rev_workbook():
    w=load_workbook(BytesIO(workbook()))
    s=w.create_sheet('Rev 관리',0)
    s.cell(2,4,'통합 이력');s.cell(5,2,'Total');s.cell(7,2,1);s.cell(7,3,1);s.cell(8,3,1)
    for col,value in {2:'000',3:'SW',4:'원본 변경내용',10:'2026.10.16',11:'원작성자'}.items():
        s.cell(12,col,value)
    b=BytesIO();w.save(b);return b.getvalue()


def test_revision_create_edit_audit_and_isolation(tmp_path):
    data=rev_workbook(); original=load_revision(data)['history']; before=original.copy(deep=True)
    db=tmp_path/'pending.sqlite3';key=source_key(data)
    existing=read_revisions(db,key,original)
    assert suggest_revision(existing)=='001'
    values={'Rev No.':'001','구분':'HW','주요 변경내용':'첫 줄\n둘째 줄','변경일자':'2026-10-22','작성자':'작성자','변경점 Key Word / 비고':'Robot'}
    save_revision(db,key,original,values,'등록자')
    revisions=read_revisions(db,key,original)
    assert len(revisions)==2 and revisions.iloc[-1]['출처']=='화면 등록'
    assert read_items(db,key).empty
    assert len(read_revisions(db,'other',original))==1
    for number in ['1','0001','000']:
        with pytest.raises(ValueError,match='이미 존재'):
            save_revision(db,key,original,{**values,'Rev No.':number},'등록자')
    for patch in [{'Rev No.':'ABC'},{'주요 변경내용':''},{'작성자':''}]:
        with pytest.raises(ValueError): save_revision(db,key,original,{**values,**patch},'등록자')
    with pytest.raises(ValueError): save_revision(db,key,original,values,'')
    current=revisions.iloc[0].to_dict()
    edits={**current,'주요 변경내용':'원본 수정 내용','Rev No.':'002'}
    assert save_revision(db,key,original,edits,'수정자',current)==2
    edited=read_revisions(db,key,original).iloc[0].to_dict()
    assert edited['출처']=='Excel 원본 · 앱 수정'
    assert edited['원본 행']=='12'
    assert edited['Rev No.']=='002'
    assert save_revision(db,key,original,edited,'수정자',edited)==0
    with pytest.raises(ValueError,match='다른 화면'):
        save_revision(db,key,original,edits,'수정자',current)
    with pytest.raises(ValueError,match='이미 존재'):
        save_revision(db,key,original,{**edited,'Rev No.':'001'},'수정자',edited)
    audit=read_revision_history(db,key)
    assert set(audit['작업'])=={'등록','수정'}
    assert len(read_revision_history(db,key,edited['ID']))==2
    assert read_revision_history(db,'other').empty
    pd.testing.assert_frame_equal(original,before)
    new_rev=read_revisions(db,key,original).iloc[-1].to_dict()
    save_revision(db,key,original,{**new_rev,'변경점 Key Word / 비고':'수정 비고'},'수정자',new_rev)
    assert read_revisions(db,key,original).iloc[-1]['변경점 Key Word / 비고']=='수정 비고'


def test_revision_empty_and_ui_restart(tmp_path):
    root=Path(__file__).resolve().parents[1]
    for name in ['app.py','pending_data.py','pending_store.py']:
        shutil.copyfile(root/name,tmp_path/name)
    (tmp_path/'data').mkdir()
    content=rev_workbook()
    source=tmp_path/'data/default.xlsx';source.write_bytes(content)
    app=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    app.radio(key='main_page').set_value('REV 관리').run()
    assert not app.exception
    app.text_area(key='new_rev_text').set_value('새 REV 화면 등록')
    app.text_input(key='new_rev_author').set_value('작성자')
    app.text_input(key='new_rev_actor').set_value('등록자')
    next(x for x in app.button if x.label=='새 REV 저장').click().run()
    assert not app.exception
    db=tmp_path/'data/pending.sqlite3';key=source_key(content);original=load_revision(content)['history']
    current=read_revisions(db,key,original).iloc[-1].to_dict()
    app.text_area(key=f"rev_edit_{current['ID']}_{current['버전']}_주요 변경내용").set_value('화면 수정 성공')
    app.text_input(key=f"rev_actor_{current['ID']}").set_value('수정자')
    next(x for x in app.button if x.label=='REV 수정 저장').click().run()
    assert not app.exception
    again=AppTest.from_file(tmp_path/'app.py',default_timeout=60).run()
    again.radio(key='main_page').set_value('REV 관리').run()
    assert not again.exception
    assert any('화면 수정 성공' in x.value for x in again.text)
    assert len(read_revisions(db,key,original))==2
    assert source.read_bytes()==content
    empty=original.iloc[0:0]
    assert suggest_revision(empty)=='000'
    assert read_revisions(db,'empty',empty).empty


def test_revision_display_cleanup_and_html():
    from pending_data import revision_display, revision_table_html
    original=pd.DataFrame([{'Rev No.':'002','구분':'H/W\nS/W','주요 변경내용':'첫 줄\n<script>위험</script>','변경일자':'2025-08-07\n\n2025-08-07','작성자':'하대오\n이계문','변경점 Key Word / 비고':'비고/n유지','원본 행':14},
                           {'Rev No.':'003','구분':'SW','주요 변경내용':'다른 내용','변경일자':'2025-08-07\n2025-08-08','작성자':'작성자','변경점 Key Word / 비고':'A\\nB','원본 행':15}])
    before=original.copy(deep=True)
    displayed,report=revision_display(original)
    assert displayed.iloc[0]['변경일자']=='2025-08-07'
    assert displayed.iloc[1]['변경일자']=='2025-08-07 / 2025-08-08'
    assert displayed.iloc[0]['작성자']=='하대오 / 이계문'
    assert displayed.iloc[0]['변경점 Key Word / 비고']=='비고/유지'
    assert displayed.iloc[1]['변경점 Key Word / 비고']=='A\\B'
    assert len(report)==5
    html=revision_table_html(displayed)
    assert 'H/W<br>S/W' in html
    assert '\\n' not in html and '/n' not in html
    assert '<script>' not in html and '&lt;script&gt;' in html
    pd.testing.assert_frame_equal(original,before)
